"""Write report tables to one Excel workbook: a contents sheet, then one sheet per table.

Each table sheet has a single header row (column name and unit) and one row per record, so
it can be filtered and pivoted directly.
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from rfnbo.report.tables import Column, Table

FORMATS = {
    "kg": "#,##0.000",
    "kWh": "#,##0.000",
    "MJ": "#,##0.000",
    "gCO2e": "#,##0.000",
    "gCO2e/MJ": "0.000000",
    "": "General",
}
HEADER = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def _format(col: Column) -> str:
    if col.kind == "integer":
        return "0"
    if col.kind == "datetime":
        return "yyyy-mm-dd hh:mm"
    if col.kind == "number" and col.unit == "" and col.key in ("rfnbo_share",):
        return "0.000000"
    return FORMATS.get(col.unit, "#,##0.000000") if col.kind == "number" else "General"


def write_xlsx(tables: Sequence[Table], path: str | Path) -> Path:
    wb = Workbook()
    contents = wb.active
    assert contents is not None
    contents.title = "Contents"
    contents.append(["Sheet", "Table", "Rows", "Note"])
    for cell in contents[1]:
        cell.font, cell.fill = HEADER, HEADER_FILL

    used: set[str] = {"Contents"}
    for table in tables:
        name = _sheet_name(table.name, used)
        ws = wb.create_sheet(name)
        ws.append([c.title for c in table.columns])
        for cell in ws[1]:
            cell.font, cell.fill = HEADER, HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in table.rows:
            ws.append([_value(row.get(c.key)) for c in table.columns])
        for j, col in enumerate(table.columns, start=1):
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(10, min(48, len(col.title) + 2))
            fmt = _format(col)
            for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                cell.number_format = fmt
        ws.freeze_panes = "B2"
        if table.rows:
            ws.auto_filter.ref = ws.dimensions
        contents.append([name, table.title, len(table.rows), table.note])
    contents.append([])
    contents.append(
        [
            "",
            "Numbers come from the calculator. Excel keeps 15 significant digits; "
            "the CSV export keeps full precision.",
        ]
    )
    for letter, width in (("A", 24), ("B", 60), ("C", 8), ("D", 100)):
        contents.column_dimensions[letter].width = width

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _sheet_name(name: str, used: set[str]) -> str:
    clean = "".join(ch for ch in name if ch not in "[]:*?/\\")[:31] or "Sheet"
    base, n = clean, 2
    while clean in used:
        suffix = f" {n}"
        clean = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(clean)
    return clean


def _value(v: object) -> object:
    if isinstance(v, dt.datetime) and v.tzinfo is not None:
        raise ValueError("Excel cannot hold time zones; pass local or UTC text")
    return v
