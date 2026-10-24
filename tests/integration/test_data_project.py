"""Phase 17: a project folder set up, checked, corrected and run through the Data screen's
state (no browser), on test case 4's files; the CLI runs the same folder."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
from pathlib import Path

import pytest

from rfnbo.cli import main
from rfnbo.data import calendar as cal
from rfnbo.project.folder import Project, prepare
from rfnbo.trace import run as sealed
from rfnbo.ui.data.workspace import Workspace, WorkspaceError

GOLDEN = Path(__file__).resolve().parents[1] / "golden"
PLANT = GOLDEN / "tc4.plant.json"
WORKBOOK = GOLDEN / "workbooks" / "RFNBO_optimiser_test4_purchased_Q4_2026.xlsx"
SOURCE = json.loads((GOLDEN / "tc4.source.json").read_text("utf-8"))


def _workspace(root: Path) -> Workspace:
    ws = Workspace(root)
    ws.setup(str(PLANT), "2026-Q4")
    i = ws.add_data_file(WORKBOOK.name, WORKBOOK.read_bytes())
    assert ws.project is not None and ws.project.file.sources[i].sheet == "Data_15min"
    assert ws.import_mapping(i, SOURCE) == len(ws.required_series())
    return ws


def test_upload_map_and_preview_in_plant_time(tmp_path: Path) -> None:
    ws = _workspace(tmp_path)
    assert all(r.mapped for r in ws.mapping())
    assert [f for f in ws.flags() if f.severity == "error"] == []
    pv = ws.preview(0, 3)
    assert pv["intervals"] == 8836
    assert pv["rows"][0]["local"] == "2026-10-01T00:00+02:00"
    assert pv["rows"][0]["utc"] == "2026-09-30T22:00:00+00:00"
    assert pv["rows"][0]["renewable_available_kWh"] == 3000.0
    # The original file is kept as uploaded.
    assert (tmp_path / "data" / WORKBOOK.name).read_bytes() == WORKBOOK.read_bytes()


def _frame_csv(tmp_path: Path, blank: range = range(0)) -> bytes:
    """Test case 4's data as a CSV whose headers are the series names, some cells blank."""
    ws = _workspace(tmp_path / "src")
    frame = ws.prepared().frame
    text = sealed.frame_csv(frame).decode()
    rows = list(csv.reader(io.StringIO(text)))
    col = rows[0].index("liquefier_electricity_kWh")
    for i in blank:
        rows[1 + i][col] = ""
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(rows)
    return out.getvalue().encode()


def test_headers_that_agree_are_mapped_automatically(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "p")
    ws.setup(str(PLANT), "2026-Q4")
    i = ws.add_data_file("plant.csv", _frame_csv(tmp_path))
    assert ws.project is not None
    src = ws.project.file.sources[i]
    assert src.time_column == "timestamp_utc"
    assert all(r.mapped for r in ws.mapping())
    assert [f for f in ws.flags() if f.severity == "error"] == []


def test_gaps_stay_errors_until_corrected_with_a_reason(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "p")
    ws.setup(str(PLANT), "2026-Q4")
    ws.add_data_file("plant.csv", _frame_csv(tmp_path, blank=range(4, 8)))
    errors = [f for f in ws.flags() if f.severity == "error"]
    assert [(f.code, f.series, f.count) for f in errors] == [
        ("missing_value", "liquefier_electricity_kWh", 4)
    ]
    assert ws.first_flagged() == 4
    with pytest.raises(WorkspaceError, match="reason"):
        ws.add_correction(
            "liquefier_electricity_kWh", "2026-10-01 01:00", "2026-10-01 02:00", 300, " "
        )
    cid = ws.add_correction(
        "liquefier_electricity_kWh",
        "2026-10-01 01:00",
        "2026-10-01 02:00",
        300.0,
        "meter offline; value from the liquefier's DCS log",
    )
    assert [f for f in ws.flags() if f.severity == "error"] == []
    (applied,) = ws.prepared().applied
    assert applied.missing_before == 4 and applied.corrected_sum == 1200
    row = ws.preview(4, 1)["rows"][0]
    assert row["liquefier_electricity_kWh"] == 300 and row["liquefier_electricity_kWh__corrected"]
    # Each value has one stated origin: overlapping corrections are refused.
    with pytest.raises(WorkspaceError, match="overlaps"):
        ws.add_correction(
            "liquefier_electricity_kWh", "2026-10-01 01:30", "2026-10-01 03:00", 1, "x"
        )
    # A wall time the clock change makes ambiguous needs its offset.
    with pytest.raises(WorkspaceError, match="occurs twice"):
        ws.add_correction(
            "liquefier_electricity_kWh", "2026-10-25 02:15", "2026-10-25 04:00", 1, "x"
        )
    ws.remove_correction(cid)
    assert any(f.severity == "error" for f in ws.flags())


def test_forms_refuse_what_the_rules_refuse(tmp_path: Path) -> None:
    ws = _workspace(tmp_path)
    ws.add_consignment(
        id="CA",
        output="consignment_a",
        volume_kg=150000,
        ship_date="2026-12-15",
        target_gCO2e_per_MJ=5.0,
    )
    with pytest.raises(WorkspaceError, match="not a consignment output"):
        ws.add_consignment(id="CX", output="lh2_storage", volume_kg=1, ship_date="2026-12-15")
    with pytest.raises(WorkspaceError, match="already exists"):
        ws.add_consignment(id="CA", output="consignment_a", volume_kg=1, ship_date="2026-12-15")
    pos = {"supplier": "H2 Trailers BV", "product": "H2", "quantity_kg": 100.0}
    ws.add_pos(id="P1", dispatch_date="2026-10-20", ghg_gCO2e_per_MJ=3.0, **pos)
    with pytest.raises(WorkspaceError, match="pos_compliant"):
        ws.add_pos(id="P2", dispatch_date="2026-10-20", ghg_gCO2e_per_MJ=28.3, **pos)
    with pytest.raises(WorkspaceError, match="pos_in_quarter"):
        ws.add_pos(id="P3", dispatch_date="2027-01-02", ghg_gCO2e_per_MJ=3.0, **pos)
    ws.set_opening(
        "lh2_storage",
        1000.0,
        [
            {
                "batch_id": "L-2026-09",
                "rfnbo_kg": 400.0,
                "ghg_gCO2e_per_MJ": 1.7,
                "production_month": "2026-09",
            }
        ],
    )
    with pytest.raises(WorkspaceError, match="not a storage"):
        ws.set_opening("pt_compressor", 0.0, [])
    cfg = ws.config()
    lh2 = cfg.block("lh2_storage")
    assert lh2.opening_ledger[0].rfnbo_kg == 400.0  # type: ignore[union-attr]
    ws.save()
    again = Project.load(tmp_path)
    assert [c.id for c in again.consignments] == ["CA"]
    assert [p.id for p in again.pos] == ["P1"]
    assert again.opening["lh2_storage"].opening_ledger[0].batch_id == "L-2026-09"
    assert again.file == ws.project.file  # type: ignore[union-attr]


def test_run_from_the_data_screen_and_from_the_cli(tmp_path: Path) -> None:
    """The screen and ``rfnbo optimise --project`` read the same files; a correction is
    recorded with the original file's hash and listed in the report."""
    ws = _workspace(tmp_path)
    decisions = json.loads((GOLDEN / "tc4.decisions.json").read_text("utf-8"))
    for c in decisions["consignments"]:
        ws.add_consignment(**c)
    for p in decisions["pos"]:
        ws.add_pos(**p)
    with pytest.raises(WorkspaceError, match="no consignments"):
        _workspace(tmp_path / "no-consignments").run()
    ws.add_correction(
        "compressor_electricity_kWh",
        "2026-10-02 00:00",
        "2026-10-02 00:30",
        40.0,
        "confirmed against the compressor's own log (no change in value)",
    )
    out = ws.run("first")
    assert out.out == tmp_path / "runs" / "first"
    rec = sealed.load_record(out.out / "optimised")
    (corr,) = rec["provenance"]["corrections"]
    assert corr["reason"].startswith("confirmed") and corr["intervals"] == 2
    assert rec["provenance"]["source_files_sha256"] == [sealed.sha256(WORKBOOK.read_bytes())]
    assert sealed.verify(out.out / "optimised") == []
    assert out.summary["objectives"]["optimised"]["consignments_met"]
    import openpyxl

    wb = openpyxl.load_workbook(out.report, read_only=True)
    assert "Data corrections" in wb.sheetnames and "PoS register" in wb.sheetnames

    assert main(["optimise", "--project", str(tmp_path), "--label", "cli"]) == 0
    cli = sealed.load_record(tmp_path / "runs" / "cli" / "optimised")
    assert cli["run_id"] == rec["run_id"]  # same inputs, same run


def test_project_files_are_plain(tmp_path: Path) -> None:
    ws = _workspace(tmp_path)
    ws.add_pos(
        id="POS-1",
        supplier="H2 Trailers BV",
        product="H2",
        dispatch_date="2026-11-01",
        quantity_kg=10.0,
        ghg_gCO2e_per_MJ=3.0,
    )
    ws.save()
    names = sorted(p.name for p in tmp_path.iterdir())
    assert names == [
        "consignments.json",
        "corrections.json",
        "data",
        "opening.json",
        "pos.csv",
        "project.json",
    ]
    assert (tmp_path / "pos.csv").read_text("utf-8").splitlines()[1] == (
        "POS-1,H2 Trailers BV,H2,2026-11-01,10.0,3.0"
    )
    assert prepare(Project.load(tmp_path)).ok


def test_local_times_convert_through_the_calendar() -> None:
    tz = "Europe/Amsterdam"
    assert cal.from_local(dt.datetime(2026, 10, 1, 0, 0), tz) == dt.datetime(
        2026, 9, 30, 22, 0, tzinfo=dt.UTC
    )
    with pytest.raises(ValueError, match="does not exist"):
        cal.from_local(dt.datetime(2027, 3, 28, 2, 30), tz)
    aware = dt.datetime.fromisoformat("2026-10-25 02:15+01:00")
    assert cal.from_local(aware, tz) == dt.datetime(2026, 10, 25, 1, 15, tzinfo=dt.UTC)


def test_status_says_what_stops_a_run(tmp_path: Path) -> None:
    """The Run tab's checklist and the tab badges: each required step, and why it is not done."""
    blank = Workspace(tmp_path / "blank")
    assert [(s.tab, s.ok) for s in blank.status()] == [("Setup", False)]

    ws = Workspace(tmp_path / "p")
    ws.setup(str(PLANT), "2026-Q4")
    steps = {s.tab: s for s in ws.status()}
    assert steps["Mapping"].badge == f"0/{len(ws.required_series())}"
    assert not steps["Preview"].ok and "map every series" in steps["Preview"].detail
    assert [s.tab for s in ws.blocking(steps.values())] == ["Mapping", "Preview", "Consignments"]

    ws.add_data_file("plant.csv", _frame_csv(tmp_path, blank=range(4, 8)))
    steps = {s.tab: s for s in ws.status()}
    assert steps["Mapping"].ok
    assert steps["Preview"].badge == "1" and "1 data quality error" in steps["Preview"].detail
    ws.add_correction(
        "liquefier_electricity_kWh", "2026-10-01 01:00", "2026-10-01 02:00", 300.0, "DCS log"
    )
    ws.add_consignment(id="CA", output="consignment_a", volume_kg=1, ship_date="2026-12-15")
    steps = {s.tab: s for s in ws.status()}
    assert steps["Preview"].ok and steps["Corrections"].badge == "1"
    assert ws.blocking(steps.values()) == []  # PoS and the opening ledger are optional


def test_records_are_edited_in_place(tmp_path: Path) -> None:
    ws = _workspace(tmp_path)
    cons = {"output": "consignment_a", "ship_date": "2026-12-15"}
    ws.add_consignment(id="CA", volume_kg=100, **cons)
    ws.add_consignment(id="CB", volume_kg=200, **cons)
    ws.add_consignment(replace="CA", id="CA2", volume_kg=150, **cons)
    assert ws.project is not None
    assert [(c.id, c.volume_kg) for c in ws.project.consignments] == [("CA2", 150), ("CB", 200)]
    with pytest.raises(WorkspaceError, match="already exists"):
        ws.add_consignment(replace="CA2", id="CB", volume_kg=1, **cons)
    with pytest.raises(WorkspaceError, match="no consignment CA"):
        ws.add_consignment(replace="CA", id="CA", volume_kg=1, **cons)

    pos = {"supplier": "H2 Trailers BV", "product": "H2", "dispatch_date": "2026-10-20"}
    ws.add_pos(id="P1", quantity_kg=100.0, ghg_gCO2e_per_MJ=3.0, **pos)
    with pytest.raises(WorkspaceError, match="pos_compliant"):
        ws.add_pos(replace="P1", id="P1", quantity_kg=100.0, ghg_gCO2e_per_MJ=28.3, **pos)
    ws.add_pos(replace="P1", id="P1", quantity_kg=120.0, ghg_gCO2e_per_MJ=3.0, **pos)
    assert [(x.id, x.quantity_kg) for x in ws.project.pos] == [("P1", 120.0)]

    first = ("compressor_electricity_kWh", "2026-10-02 00:00", "2026-10-02 00:30", 40.0, "log")
    cid = ws.add_correction(*first)
    ws.add_correction(*first[:3], 41.0, "log, re-read", replace=cid)
    (c,) = ws.project.corrections
    assert (c.id, c.value, c.reason) == (cid, 41.0, "log, re-read")
    # Shown in plant time with its offset, which the form takes back unchanged (and which
    # tells the two 02:00–03:00 hours apart on the night the clocks go back).
    assert ws.local(c.start) == "2026-10-02 00:00+02:00"

    ws.set_opening(
        "lh2_storage",
        1000.0,
        [
            {"batch_id": b, "rfnbo_kg": 1.0, "ghg_gCO2e_per_MJ": 1.7, "production_month": "2026-09"}
            for b in ("L1", "L2")
        ],
    )
    ws.remove_opening_entry("lh2_storage", "L1")
    o = ws.project.opening["lh2_storage"]
    assert o.opening_inventory_kg == 1000.0
    assert [e.batch_id for e in o.opening_ledger] == ["L2"]
    with pytest.raises(WorkspaceError, match="no opening entry"):
        ws.remove_opening_entry("lh2_storage", "L1")


def test_plant_config_is_chosen_from_the_folder_or_uploaded(tmp_path: Path) -> None:
    ws = Workspace(tmp_path / "p")
    assert ws.plant_configs() == []  # the folder need not exist yet
    name = ws.add_plant_config("tc4.plant.json", PLANT.read_bytes())
    assert (tmp_path / "p" / name).read_bytes() == PLANT.read_bytes()
    assert ws.add_plant_config(name, PLANT.read_bytes()) == name  # the same file again
    with pytest.raises(WorkspaceError, match="different file"):
        ws.add_plant_config(name, b'{"plant": "other"}')
    with pytest.raises(WorkspaceError, match="project's own files"):
        ws.add_plant_config("project.json", b"{}")
    with pytest.raises(WorkspaceError, match="not a plant config"):
        ws.add_plant_config("x.json", b"[1]")
    # A draft is kept (it opens in the editor) but cannot be used.
    ws.add_plant_config("draft.json", b'{"plant": "half drawn"}')
    assert ws.plant_configs() == ["draft.json", "tc4.plant.json"]
    draft = ws.plant_summary("draft.json")
    assert not draft.ok and draft.issues
    assert not ws.plant_summary("missing.json").ok
    good = ws.plant_summary(name)
    assert good.ok and good.plant is not None and good.plant.timezone == "Europe/Amsterdam"
    assert ws.new_plant_name() == "plant.json"

    quarters, suggested = ws.quarters(good.plant, dt.date(2026, 10, 24))
    assert quarters == ["2026-Q4", "2026-Q3", "2026-Q2", "2026-Q1"]  # from effective_from
    assert suggested == "2026-Q3"  # the last complete quarter
    ws.setup(name, "2026-Q4")
    ws.save()
    assert Project.load(tmp_path / "p").file.config == "tc4.plant.json"  # inside the folder
    assert ws.plant_configs() == ["draft.json", "tc4.plant.json"]
