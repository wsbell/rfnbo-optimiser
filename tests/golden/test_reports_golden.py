"""Phase 8 acceptance: reports regenerate the workbook tables.

Test case 1 is compared with the workbook's own sheets, cell by cell, except where the
no-averaging rule (§12) changes the workbook's values (NH₃ GHG, the NH₃ ledger and the
consignments): those are compared with the regenerated fixture. Test case 2 is compared with its
regenerated fixture (the workbook predates the §2.5 and §12 rules).
"""

from __future__ import annotations

import csv
import math
from functools import cache
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from rfnbo.config.schema import ProcessingUnit
from rfnbo.report import tables as T
from rfnbo.report.csv_export import file_name, write_csv
from rfnbo.report.excel import write_xlsx
from tests.golden.test_golden import CASES, _run

GOLDEN = Path(__file__).resolve().parent
GRID = "em:grid_power|Non-renewable electricity"
WATER = "em:water|Water"


@cache
def _sheet(case: str, name: str) -> list[tuple[Any, ...]]:
    wb = openpyxl.load_workbook(GOLDEN / "workbooks" / CASES[case], data_only=True, read_only=True)
    rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    return rows


def _labelled(case: str, name: str) -> dict[str, Any]:
    return {str(r[0]).strip(): r[1] for r in _sheet(case, name) if r and r[0] is not None}


def same(a: Any, b: Any, what: str) -> None:
    if isinstance(b, bool) or isinstance(a, bool):
        assert bool(a) is bool(b), what
    elif isinstance(b, int | float) and not isinstance(b, bool):
        assert a is not None and math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-6), (
            f"{what}: report {a!r}, workbook {b!r}"
        )
    else:
        assert a == b, f"{what}: report {a!r}, workbook {b!r}"


def _table(case: str, name: str) -> T.Table:
    run, _, _ = _run(case)
    return next(t for t in T.all_tables(run) if t.name == name)


def _by(table: T.Table, key: str = "batch") -> dict[Any, dict[str, Any]]:
    return {r[key]: r for r in table.rows}


def test_tc1_electrolyser_register_is_the_h2_batches_sheet() -> None:
    rows = _by(_table("tc1", "Batches Electrolyser"))
    keys = [
        "batch",
        "month",
        "start_local",
        "end_local",
        "intervals",
        "renewable_available_kWh",
        "electricity_kWh",
        "renewable_matched_kWh",
        "rfnbo_share",
        "product_kg",
        "energy_MJ",
        WATER,
        GRID,
        "ghg",
        "compliant",
        "rfnbo_kg",
        "rfnbo_MJ",
    ]
    wb_rows = [r for r in _sheet("tc1", "H2_Batches")[1:] if str(r[0]).startswith("E-")]
    assert set(rows) == {r[0] for r in wb_rows}
    for r in wb_rows:
        for key, value in zip(keys, r, strict=False):
            same(rows[r[0]][key], value, f"{r[0]} {key}")


def test_tc1_nh3_register_is_the_nh3_batches_sheet() -> None:
    rows = _by(_table("tc1", "Batches NH3 synthesis"))
    cols = {
        "month": 1,
        "intervals": 2,
        "product_kg": 3,
        "feed_consumed_kg": 4,
        "rfnbo_feed_kg": 5,
        "feed_ghg": 7,
        "rfnbo_share": 8,
        "feed_emissions_g": 9,
        "electricity_kWh": 10,
        GRID: 11,
        "rfnbo_kg": 12,
        "rfnbo_MJ": 13,
    }
    _, expected, _ = _run("tc1")
    wb_rows = [
        r
        for r in _sheet("tc1", "NH3_Batches")[1:]
        if str(r[0]).startswith("N-") and "opening" not in str(r[0])
    ]
    assert set(rows) == {r[0] for r in wb_rows}
    for r in wb_rows:
        for key, j in cols.items():
            same(rows[r[0]][key], r[j], f"{r[0]} {key}")
        e = expected["nh3_batches"][r[0]]
        assert rows[r[0]]["ghg"] == e["ghg"] or math.isclose(rows[r[0]]["ghg"], e["ghg"]), r[0]
        same(rows[r[0]]["compliant"], e["compliant"], f"{r[0]} compliant")


def test_tc1_feed_matrix_is_the_nh3_feed_sheet() -> None:
    table = _table("tc1", "Feed NH3 synthesis")
    rows = _by(table)
    head = _sheet("tc1", "NH3_Feed")[0]
    targets = {j: str(h).split()[1] for j, h in enumerate(head) if str(h).startswith("→ N-")}
    for r in _sheet("tc1", "NH3_Feed")[1:]:
        if not str(r[0]).startswith("E-") or not r[2]:
            continue  # 0% batches hold no RFNBO and are not in the ledger
        row = rows[r[0]]
        same(row["ghg"], r[1], f"{r[0]} GHG")
        same(row["available_kg"], r[2], f"{r[0]} available")
        for j, target in targets.items():
            same(row[f"to:{target}"], r[j], f"{r[0]} → {target}")
        same(row["left_kg"], r[6], f"{r[0]} left")


def test_tc1_ledgers_and_reconciliation_are_the_ledger_sheets() -> None:
    lines = {(r["storage"], r["batch"]): r for r in _table("tc1", "Ledger batches").rows}
    _, expected, _ = _run("tc1")
    nh3_lines = expected["nh3_ledger"]["lines"]
    assert {b for s, b in lines if s == "nh3_storage"} == set(nh3_lines)
    for bid, x in nh3_lines.items():
        for key in (
            "month",
            "ghg",
            "in_kg",
            "out_consignment_kg",
            "closing_before_kg",
            "write_off_kg",
            "carried_kg",
        ):
            same(lines["nh3_storage", bid][key], x[key], f"NH3 ledger {bid} {key}")

    summary = {r["storage"]: r for r in _table("tc1", "Ledgers").rows}
    h2 = _labelled("tc1", "H2_Ledger")
    for key, label in (
        ("opening_kg", "Opening RFNBO ledger (kg)"),
        ("produced_kg", "RFNBO H₂ produced in quarter (kg)"),
        ("out_feed_kg", "RFNBO H₂ allocated to NH₃ synthesis (kg)"),
        ("closing_before_kg", "Closing RFNBO ledger before write-off (kg)"),
        ("physical_kg", "Measured physical inventory at quarter end (kg)"),
        ("write_off_kg", "Write-off (kg)"),
        ("carried_kg", "RFNBO H₂ carried to Q1 2027 (kg)"),
        ("carried_ghg", "Energy-weighted GHG of carried H₂ (gCO₂e/MJ)"),
    ):
        same(summary["h2_storage"][key], h2[label], f"H2 ledger {key}")
    nh3 = _labelled("tc1", "NH3_Ledger")
    same(
        summary["nh3_storage"]["physical_kg"],
        nh3["Measured physical inventory at quarter end (kg)"],
        "NH3 physical",
    )
    same(summary["nh3_storage"]["write_off_kg"], expected["nh3_ledger"]["write_off_kg"], "NH3 w/o")
    same(summary["nh3_storage"]["carried_ghg"], expected["nh3_ledger"]["carried_ghg"], "NH3 GHG")

    rec = {(r["storage"], r["item"]): r["kg"] for r in _table("tc1", "Reconciliation").rows}
    for store, sheet, items in (
        (
            "h2_storage",
            "H2_Ledger",
            (
                ("Opening physical inventory", "Opening physical inventory (kg)"),
                ("Produced into storage", "H₂ produced (kg)"),
                ("Out to NH3 synthesis (f_h2s_nh3)", "H₂ consumed by NH₃ synthesis (kg)"),
                ("Expected closing", "Expected closing (kg)"),
                ("Measured closing", "Measured closing (kg)"),
                ("Implied loss", "Implied loss (kg)"),
            ),
        ),
        (
            "nh3_storage",
            "NH3_Ledger",
            (
                ("Opening physical inventory", "Opening physical inventory (kg)"),
                ("Produced into storage", "NH₃ produced (kg)"),
                ("Out to Distribution to Customer A (f_nh3s_a)", "Shipped C1 (kg)"),
                ("Out to Distribution to Customer B (f_nh3s_b)", "Shipped C2 (kg)"),
                ("Out to Spot sales (no RFNBO claim) (f_nh3s_spot)", "Spot sale, no claim (kg)"),
                ("Expected closing", "Expected closing (kg)"),
                ("Measured closing", "Measured closing (kg)"),
                ("Implied loss", "Implied loss (kg)"),
            ),
        ),
    ):
        labels = _labelled("tc1", sheet)
        for item, label in items:
            same(rec[store, item], labels[label], f"{store} {item}")


def test_tc1_consignment_report_is_the_consignments_sheet() -> None:
    _, expected, _ = _run("tc1")
    summary = _by(_table("tc1", "Consignments"), "consignment")
    parts = {
        (r["consignment"], r["month"]): r["kg"] for r in _table("tc1", "Consignment batches").rows
    }
    for r in _sheet("tc1", "Consignments")[1:]:
        if r[0] not in ("C1", "C2"):
            continue
        row, e = summary[r[0]], expected["consignments"][r[0]]
        same(row["customer"], r[1], "customer")
        same(row["ship_date"], r[2].date().isoformat(), "ship date")
        for key, j in (
            ("volume_kg", 3),
            ("target", 4),
            ("adder", 5),
        ):
            same(row[key], r[j], f"{r[0]} {key}")
        # §12: the workbook's averaged allocation is superseded by the regenerated one.
        same(row["highest_ghg"], e["highest_ghg"], f"{r[0]} highest part GHG")
        same(row["shipped_ghg"], e["shipped_ghg"], f"{r[0]} shipped GHG")
        same(row["meets_target"], e["shipped_ghg"] <= e["target"] + 1e-9, f"{r[0]} target")
        for month in ("2026-09", "2026-10"):
            kg = sum(v for b, v in e["parts"].items() if b.startswith(f"N-{month}"))
            same(parts.get((r[0], month), 0.0), kg, f"{r[0]} {month} NH3")


def test_tc1_renewable_matching_totals() -> None:
    rows = _by(_table("tc1", "Renewable matching"), "period")
    h2 = [r for r in _sheet("tc1", "H2_Batches")[1:] if str(r[0]).startswith("E-")]
    for month in ("2026-10", "2026-11", "2026-12"):
        wb = [r for r in h2 if r[1] == month]
        same(rows[month]["available_kWh"], sum(r[5] for r in wb), f"{month} available")
        same(rows[month]["unit:electrolyser"], sum(r[7] for r in wb), f"{month} matched")


def test_tc2_registers_and_ledgers_match_the_regenerated_fixture() -> None:
    _, expected, _ = _run("tc2")
    h2 = _by(_table("tc2", "Batches Electrolyser"))
    for bid, e in expected["h2_batches"].items():
        for key, ekey in (
            ("rfnbo_share", "rfnbo_share"),
            ("ghg", "ghg"),
            ("rfnbo_kg", "rfnbo_kg"),
            (GRID, "grid_emissions_g"),
            (WATER, "water_emissions_g"),
            ("intervals", "intervals"),
        ):
            same(h2[bid][key], e[ekey], f"{bid} {key}")
    nh3 = _by(_table("tc2", "Batches NH3 synthesis"))
    for bid, e in expected["nh3_batches"].items():
        for key in ("rfnbo_share", "ghg", "rfnbo_kg", "feed_ghg", "rfnbo_feed_kg"):
            same(nh3[bid][key], e[key], f"{bid} {key}")
    summary = {r["storage"]: r for r in _table("tc2", "Ledgers").rows}
    for key in ("produced_kg", "closing_before_kg", "write_off_kg", "carried_kg", "carried_ghg"):
        same(summary["h2_storage"][key], expected["h2_ledger"][key], f"H2 ledger {key}")


def test_calculation_log_lists_every_check_with_its_reference() -> None:
    run, _, _ = _run("tc2")
    log = _table("tc2", "Calculation log")
    checks = [r for r in log.rows if r["kind"] == "rule check"]
    assert len(checks) == len(run.checks)
    assert all(r["reference"] and r["result"] == "pass" for r in checks)


@pytest.mark.parametrize("case", sorted(CASES))
def test_excel_export_round_trips(case: str, tmp_path: Path) -> None:
    run, _, _ = _run(case)
    tables = T.all_tables(run, "run-test")
    path = write_xlsx(tables, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    contents = [r[0] for r in wb["Contents"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(contents) == len(tables)
    for name, table in zip(contents, tables, strict=True):
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        assert list(rows[0]) == [c.title for c in table.columns]
        assert len(rows) - 1 == len(table.rows)
        for got, want in zip(rows[1:], table.rows, strict=True):
            for value, col in zip(got, table.columns, strict=True):
                expected = want.get(col.key)
                if isinstance(expected, float):
                    # Excel holds 15 significant digits; the CSV export keeps all 17.
                    assert math.isclose(value, expected, rel_tol=1e-14, abs_tol=1e-300), (
                        name,
                        col.key,
                    )
                else:  # Excel reads an empty text cell back as empty
                    assert value == (None if expected == "" else expected), (name, col.key)


def test_csv_export_round_trips(tmp_path: Path) -> None:
    run, _, _ = _run("tc1")
    tables = T.all_tables(run)
    paths = write_csv(tables, tmp_path)
    assert [p.name for p in paths] == [file_name(t) for t in tables]
    register = next(t for t in tables if t.name == "Batches Electrolyser")
    with (tmp_path / file_name(register)).open(encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == [c.title for c in register.columns]
    ghg = [c.key for c in register.columns].index("ghg")
    assert [float(r[ghg]) for r in rows[1:]] == register.column("ghg")


def test_every_unit_gets_a_register() -> None:
    run, _, _ = _run("tc1")
    names = {t.name for t in T.all_tables(run)}
    for unit in run.config.blocks_of(ProcessingUnit):
        assert f"Batches {unit.name}"[:31] in names


def test_cli_report_from_a_sealed_run(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    from rfnbo.cli import main
    from rfnbo.trace import run as sealed

    run, _, flags = _run("tc1")
    d = sealed.seal(run, tmp_path / "run", flags)
    assert main(["report", str(d), "--out", str(tmp_path / "out"), "--csv"]) == 0
    xlsx = next((tmp_path / "out").glob("report_run-*.xlsx"))
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    assert "Batches Electrolyser" in wb.sheetnames and "Calculation log" in wb.sheetnames
    assert len(list((tmp_path / "out").glob("csv_run-*/*.csv"))) == len(T.all_tables(run))

    frame = d / sealed.FRAME_FILE
    frame.write_text(
        frame.read_text(encoding="utf-8").replace(",2500.0,", ",2499.0,", 1), encoding="utf-8"
    )
    assert main(["report", str(d), "--out", str(tmp_path / "bad")]) == 1
    assert "refusing" in capsys.readouterr().err


def test_batch_parts_table_lists_every_downstream_part() -> None:
    """§12: December NH₃ in test case 1 has a compliant part (from the 1.17 H₂) and a
    non-compliant one (from the 17.83 H₂); the register shows the highest."""
    _, expected, _ = _run("tc1")
    rows = _by(_table("tc1", "Batch parts"), "part")
    want = {pid: x for b in expected["nh3_batches"].values() for pid, x in b["parts"].items()}
    assert set(rows) == set(want)
    for pid, x in want.items():
        same(rows[pid]["ghg"], x["ghg"], f"{pid} GHG")
        same(rows[pid]["kg"], x["kg"], f"{pid} kg")
        same(rows[pid]["compliant"], x["compliant"], f"{pid} compliant")
    dec = _by(_table("tc1", "Batches NH3 synthesis"))["N-2026-12"]
    assert dec["parts"] == 2 and dec["ghg"] is None and not dec["compliant"]
    same(dec["max_ghg"], want["N-2026-12.2"]["ghg"], "highest part")
