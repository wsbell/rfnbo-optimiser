"""Phase 10 acceptance: test case 3 (multi-step plant) is specified before any network code.

The workbook's Checks sheet passes against an independent script (``tc3_reference``), the
reference reproduces the workbook from the workbook's own inputs, and the fixtures exported
from the workbook are current and self-consistent. The calculator and optimiser tests for this
case arrive with phases 11–15.
"""

from __future__ import annotations

import datetime as dt
import json
import math
from functools import cache
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from rfnbo.config import load_config
from rfnbo.data import ingest, quality
from tests.golden import export_tc3
from tests.golden import tc3_reference as R

GOLDEN = Path(__file__).resolve().parent
WORKBOOK = export_tc3.WORKBOOK
SERIES = {  # workbook series name -> reference series name
    "renewable_available_kWh": "renewable_kWh",
    "electrolyser_electricity_kWh": "electrolyser_kWh",
    "electrolyser_water_kg": "water_kg",
    "electrolyser_H2_kg": "h2_kg",
    "electrolyser_to_buffer_kg": "to_buffer_kg",
    "electrolyser_to_compressor_kg": "to_compressor_kg",
    "liquefier_electricity_kWh": "liquefier_kWh",
    "liquefier_H2_kg": "liquefier_h2_kg",
    "liquefier_LH2_kg": "lh2_kg",
    "compressor_electricity_kWh": "compressor_kWh",
    "compressor_H2_kg": "compressor_h2_kg",
    "compressor_CH2_kg": "ch2_kg",
    "h2_buffer_inventory_kg": "buffer_inv_kg",
    "lh2_shipped_ca_kg": "ca_kg",
    "lh2_shipped_spot_kg": "spot_kg",
    "lh2_storage_inventory_kg": "lh2_inv_kg",
    "ch2_shipped_cb_kg": "cb_kg",
    "ch2_storage_inventory_kg": "ch2_inv_kg",
}
FLOWS = {"f_el_pt": "to_compressor", "f_el_buf": "to_buffer"}


def close(a: float, b: float) -> bool:
    return math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-6)


@cache
def _fixtures() -> tuple[dict[str, Any], dict[str, Any]]:
    decisions = json.loads((GOLDEN / "tc3.decisions.json").read_text(encoding="utf-8"))
    expected = json.loads((GOLDEN / "tc3.expected.json").read_text(encoding="utf-8"))
    return decisions, expected


@cache
def _workbook_data() -> R.Data:
    """The workbook's inputs, read through the app's ingestion and aligned to the quarter."""
    source = ingest.load_source_spec(GOLDEN / "tc3.source.json")
    raw = ingest.load(WORKBOOK, source)
    frame, flags = quality.align(raw, 2026, 4, "Europe/Amsterdam")
    assert not quality.has_errors(flags), flags
    utc = [t.astimezone(dt.UTC) for t in frame.index]
    local = [t.astimezone(R.TZ) for t in utc]
    series = {ref: [float(x) for x in frame.series[name]] for name, ref in SERIES.items()}
    return R.Data(utc, local, [f"{t.year:04d}-{t.month:02d}" for t in local], series)


def _decisions(data: R.Data, fx: dict[str, Any]) -> R.Decisions:
    """The exported decision fixture as reference decisions."""
    pos = {t: i for i, t in enumerate(data.utc)}
    end = data.utc[-1] + R.STEP
    batches: dict[str, list[tuple[str, int, int]]] = {u: [] for u in R.UNITS}
    for b in fx["batches"]:
        start, stop = (dt.datetime.fromisoformat(b[k]) for k in ("start", "end"))
        batches[b["unit"]].append((b["id"], pos[start], len(data) if stop == end else pos[stop]))
    targets = {c["id"]: c["target_gCO2e_per_MJ"] for c in fx["consignments"]}
    parts: dict[str, dict[str, float]] = {}
    for a in fx["aggregated_batches"]:
        for p in a["parts"]:
            parts.setdefault(a["consignment"], {})[p["batch"]] = p["kg"]
    return R.Decisions(
        batches=batches,
        renewable={r["unit"]: [float(x) for x in r["kWh"]] for r in fx["renewable_allocations"]},
        outputs={(o["batch"], FLOWS[o["flow"]]): o["kg"] for o in fx["output_allocations"]},
        feed={(f["source"], f["target"]): f["kg"] for f in fx["feed"]},
        parts=parts,
        targets=targets,
    )


def _assert_matches(result: dict[str, Any], expected: dict[str, Any]) -> None:
    units = result["units"]
    for unit, key in (
        ("electrolyser", "electrolyser_batches"),
        ("liquefier", "liquefier_batches"),
        ("compressor", "compressor_batches"),
    ):
        assert set(units[unit]) == set(expected[key])
        for bid, want in expected[key].items():
            got = units[unit][bid]
            for k in ("rfnbo_share", "rfnbo_kg", "ghg", "intervals", "product_kg"):
                assert close(got[k], want[k]), (bid, k, got[k], want[k])
            assert got["compliant"] == want["compliant"]
    for bid, want in expected["output_split"].items():
        got = units["electrolyser"][bid]["assigned_kg"]
        assert close(got["to_compressor"], want["assigned_to_compressor_kg"])
        assert close(got["to_buffer"], want["assigned_to_buffer_kg"])
    for storage in ("h2_buffer", "pt_compressor"):
        got, want = result["ledgers"][storage], expected["ledgers"][storage]
        for k_got, k_want in (
            ("in_kg", "in_kg"),
            ("out_kg", "out_feed_kg"),
            ("closing_before_kg", "closing_before_kg"),
            ("write_off_kg", "write_off_kg"),
            ("carried_kg", "carried_kg"),
        ):
            assert close(got[k_got], want[k_want]), (storage, k_got)
    for storage in ("lh2_storage", "ch2_storage"):
        got, want = result["ledgers"][storage], expected["ledgers"][storage]
        assert close(got["physical_kg"], want["measured_kg"])
        for bid, line in want["lines"].items():
            for k in ("in_kg", "write_off_kg", "carried_kg"):
                assert close(got["lines"][bid][k], line[k]), (storage, bid, k)
    for cid, want in expected["consignments"].items():
        got = result["consignments"][cid]
        assert close(got["shipped_ghg"], want["shipped_ghg"])
        for month, kg in want["parts_by_month"].items():
            assert close(got["by_month"].get(month, 0.0), kg), (cid, month)
    got, want = result["objectives"], expected["objectives"]
    assert got["consignments_met"] == want["consignments_met"]
    assert close(got["stored_MJ"], want["stored_MJ"])
    assert close(got["carried_emissions_g"], want["carried_emissions_g"])
    assert got["batches"] == want["batches"]
    for storage, mj in want["stored_MJ_by_storage"].items():
        assert close(got["stored_MJ_by_storage"][storage], mj), storage


def test_workbook_checks_sheet_passes() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    rows = [r for r in wb["Checks"].iter_rows(min_row=2, values_only=True) if r[1] is not None]
    assert len(rows) > 90
    assert all(r[4] is True for r in rows), [r[0] for r in rows if r[4] is not True]
    flags = {r[0]: r[4] for r in wb["Checks"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert flags["All checks pass?"] is True
    assert flags["Renewable pool never over-allocated in any interval?"] is True


def test_workbook_inputs_are_the_reference_quarter() -> None:
    data, ref = _workbook_data(), R.build_data()
    assert len(data) == len(ref) == 8836
    assert data.utc == ref.utc
    for name, xs in ref.s.items():
        assert all(close(a, b) for a, b in zip(data.s[name], xs, strict=True)), name


def test_reference_reproduces_the_workbook_from_its_own_inputs() -> None:
    data = _workbook_data()
    _, expected = _fixtures()
    _assert_matches(R.run(data, R.base_decisions(data)), expected)


def test_decision_fixture_reproduces_the_expected_values() -> None:
    """The exported decisions fully determine the expected values (phase 12 runs the app on
    the same pair)."""
    data = _workbook_data()
    decisions, expected = _fixtures()
    d = _decisions(data, decisions)
    assert R.pool_ok(data, d)
    result = R.run(data, d)
    assert R.splits_ok(result)
    _assert_matches(result, expected)
    write_offs = {(w["storage"], w["batch"]): w["kg"] for w in decisions["write_offs"]}
    computed = {
        (s, b): line["write_off_kg"]
        for s, led in result["ledgers"].items()
        for b, line in led["lines"].items()
        if line["write_off_kg"] > 1e-9
    }
    assert computed.keys() == write_offs.keys()
    assert all(close(computed[k], write_offs[k]) for k in computed)


def test_fixtures_are_current() -> None:
    """Regenerate with ``uv run python tests/golden/export_tc3.py``."""
    decisions, expected = export_tc3.export(WORKBOOK)
    assert (decisions, expected) == _fixtures()


def test_hand_calculated_values() -> None:
    """Spot values worked by hand (README of the workbook)."""
    _, x = _fixtures()
    f = 2.0 * 500 / (50 * 120) + 1.0  # 100% electrolyser H₂: water and adder only
    assert close(x["electrolyser_batches"]["E-2026-11"]["ghg"], f)
    # Boil-off raises GHG: 3% in November, 2% otherwise.
    assert close(x["liquefier_batches"]["L-2026-11"]["ghg"], f * 30 / 29.1 + 0.5)
    assert close(x["liquefier_batches"]["L-2026-10"]["ghg"], f * 30 / 29.4 + 0.5)
    assert close(x["compressor_batches"]["C-2026-10"]["ghg"], f + 0.5)
    # The buffer ends full of RFNBO: 2,000 kg opening + 96 trip intervals × 30 kg.
    assert close(x["ledgers"]["h2_buffer"]["carried_kg"], 2000 + 96 * 30)
    assert close(x["liquefier_batches"]["L-2026-12"]["rfnbo_feed_kg"], 2880 * 30 - 2000)
    assert x["ledgers"]["pt_compressor"]["carried_kg"] == 0
    # LH₂: 254,132 RFNBO − 150,000 CA − 47,092 measured = 57,040 written off, December first.
    lh2 = x["ledgers"]["lh2_storage"]
    assert close(lh2["shortfall_kg"], 57040)
    assert close(lh2["lines"]["L-2026-12"]["write_off_kg"], 57040)
    assert close(x["objectives"]["stored_MJ"], (4880 + 47092 + 56720) * 120)
    assert x["objectives"]["at_upper_bound"] is True
    assert x["workbook_checks_pass"] is True


def test_variants() -> None:
    _, x = _fixtures()
    cases = R.all_cases()
    v1 = x["variants"]["v1_hourly"]
    assert v1["electrolyser_batches"] == len(cases["v1_hourly"]["units"]["electrolyser"]) == 2209
    assert v1["zero_share_batches"] == 12
    assert close(v1["stored_MJ"], cases["v1_hourly"]["objectives"]["stored_MJ"])
    v2 = x["variants"]["v2_shortfall"]
    a, b = v2["reading_a_noncompliant_counts"], v2["reading_b_compliant_only"]
    assert close(a["stored_MJ"], cases["v2_including"]["objectives"]["stored_MJ"])
    assert close(b["stored_MJ"], cases["v2_compliant"]["objectives"]["stored_compliant_MJ"])
    assert a["december_lh2_ghg"] > R.THRESHOLD >= b["december_lh2_ghg"]
    assert b["liquefier_renewable_kWh"] >= v2["least_liquefier_renewable_kWh"]
    # Each reading's answer beats the other's under its own measure.
    assert a["stored_MJ"] > b["stored_MJ"] > a["stored_compliant_MJ"]
    v3 = x["variants"]["v3_infeasible"]
    assert v3["ca_feasible"] is False and v3["cb_met"] is True
    assert v3["ca_max_volume_kg"] == 0 and v3["consignments_met"] == 1
    assert close(
        v3["lh2_write_off_kg"], cases["v3_infeasible"]["ledgers"]["lh2_storage"]["write_off_kg"]
    )


@pytest.mark.xfail(
    strict=True, reason="phase 11: pass-through storages, split outputs, unit → unit"
)
def test_plant_config_validates() -> None:
    load_config(GOLDEN / "tc3.plant.json")
