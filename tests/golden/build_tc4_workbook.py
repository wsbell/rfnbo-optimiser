"""Build the test case 4 workbook (purchased RFNBO with proofs of sustainability, build plan
§11.1, phase 16).

    uv run python tests/golden/build_tc4_workbook.py [out.xlsx]

Test plant 3 plus a supplier of H₂ delivering into the compressor and the liquefier, and three
PoS. Writes the inputs, the expected decisions and live formulas for every result, and a Checks
sheet comparing each result with ``tc4_reference`` (an independent Python calculation). The
file has no cached values until it is recalculated, for example with LibreOffice; the committed
workbook is recalculated, and ``test_tc4_workbook.py`` reads its cached values.

Sheets shared with test case 3 (E_Batches, E_Split, LH2_Ledger, CH2_Ledger) are built by
``build_tc3_workbook``'s functions on the same column layout; Data_15min adds the supplier's
metered deliveries in columns AC and AD.
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from tests.golden import build_tc3_workbook as W3
from tests.golden import tc3_reference as R
from tests.golden import tc4_reference as T
from tests.golden.build_tc3_workbook import DECISION, FONT, KEY, P, Sheet, last, rng

OUT = Path(__file__).parent / "workbooks" / "RFNBO_optimiser_test4_purchased_Q4_2026.xlsx"
DATA_COLS = [
    *W3.DATA_COLS,
    ("AC", "Supplier H₂ delivered to liquefier (kg)", "sup_liquefier_kg"),
    ("AD", "Supplier H₂ delivered to compressor (kg)", "sup_compressor_kg"),
]
PARAMS = [*W3.PARAMS, (17, "Purchased H₂ lower heating value", 120, "MJ/kg", "Supplier block")]
POS_ROWS = {"POS-A": 2, "POS-B": 3, "POS-C": 4}
# PoS_Alloc rows: (row, PoS, flow, target batch, formula). Lowest-GHG credentials first, each
# nearest its dispatch month; the delivery table is in rows 10-11. POS-C makes no compliant
# part (Parts!P13:P14), so it gets none (§12: no blend can save it).
ALLOC = [
    (2, "POS-A", "f_sup_comp", "C-2026-11", "=MIN(PoS!$E$2,$D$10)"),
    (3, "POS-B", "f_sup_comp", "C-2026-11", "=MIN(PoS!$E$3,$D$10-E2)"),
    (4, "POS-B", "f_sup_liq", "L-2026-12", "=MIN(PoS!$E$3-E3,$D$11)"),
    (5, "POS-C", "f_sup_liq", "L-2026-12", "=IF(Parts!P14,MIN(PoS!$E$4,$D$11-E4),0)"),
]
# Parts sheet rows (§12): (row, part, unit, batch, origin, source, source GHG, RFNBO feed kg,
# batch row on the unit's batch sheet).
PARTS = [
    (
        2,
        "L-2026-10",
        "liquefier",
        "L-2026-10",
        "own",
        "E-2026-10",
        "Buffer_Ledger!B2",
        "Buffer_Ledger!D2",
        2,
    ),
    (
        3,
        "L-2026-11",
        "liquefier",
        "L-2026-11",
        "own",
        "E-2026-11",
        "Buffer_Ledger!B3",
        "Buffer_Ledger!E3",
        3,
    ),
    (
        4,
        "L-2026-12.1",
        "liquefier",
        "L-2026-12",
        "own",
        "E-2026-12",
        "Buffer_Ledger!B4",
        "Buffer_Ledger!F4",
        4,
    ),
    (5, "L-2026-12.2", "liquefier", "L-2026-12", "purchased", "POS-B", "PoS!F3", "PoS_Alloc!E4", 4),
    (
        6,
        "C-2026-10",
        "compressor",
        "C-2026-10",
        "own",
        "E-2026-10",
        "PT_Ledger!B2",
        "PT_Ledger!D2",
        2,
    ),
    (
        7,
        "C-2026-11.1",
        "compressor",
        "C-2026-11",
        "own",
        "E-2026-11",
        "PT_Ledger!B3",
        "PT_Ledger!E3",
        3,
    ),
    (
        8,
        "C-2026-11.2",
        "compressor",
        "C-2026-11",
        "purchased",
        "POS-A",
        "PoS!F2",
        "PoS_Alloc!E2",
        3,
    ),
    (
        9,
        "C-2026-11.3",
        "compressor",
        "C-2026-11",
        "purchased",
        "POS-B",
        "PoS!F3",
        "PoS_Alloc!E3",
        3,
    ),
    (
        10,
        "C-2026-12",
        "compressor",
        "C-2026-12",
        "own",
        "E-2026-12",
        "PT_Ledger!B4",
        "PT_Ledger!F4",
        4,
    ),
]
PART_ROW = {p[1]: p[0] for p in PARTS}
PA, PC, PO, PL = (f"Parts!${c}$2:${c}$10" for c in "ACOL")
# Product ledgers: lines in write-off priority order (highest GHG, then newest month).
LEDGER_ORDER = {
    "LH2_Ledger": ["L-2026-12.2", "L-2026-11", "L-2026-12.1", "L-2026-10"],
    "CH2_Ledger": ["C-2026-11.3", "C-2026-11.2", "C-2026-12", "C-2026-11.1", "C-2026-10"],
}


def total_row(sheet: str) -> int:
    """The Total row of a product ledger; labelled rows follow it."""
    return 2 + len(LEDGER_ORDER[sheet])


def readme(wb: Workbook) -> None:
    s = Sheet(wb, "README", {"A": 160})
    lines = [
        (
            "RFNBO optimiser — test case 4, purchased RFNBO with proofs of sustainability, Q4 2026",
            True,
        ),
        (
            "Purpose: expected results for build plan phase 16 (§11.1). Every result cell is a "
            "live formula you can trace back to the 15-minute data and the PoS register. Column "
            "C of the Checks sheet comes from tests/golden/tc4_reference.py, a separate Python "
            "calculation that does not read these formulas.",
            False,
        ),
        ("", False),
        ("Colour key", True),
        (
            "Blue text = input value. Black text = formula. Green text = link to another sheet. "
            "Green fill = expected decision chosen by the optimiser. Yellow fill = key answer to "
            "compare against the app.",
            False,
        ),
        ("", False),
        ("Plant", True),
        (
            "Test plant 3 (test case 3's workbook explains it) plus H2 Trailers BV, a supplier "
            "of H₂ by tube trailer. Its deliveries are metered into the compressor (10 kg per "
            "interval, 1–10 November: 9,600 kg) and into the liquefier (10 kg per interval, "
            "11–20 December: 9,600 kg). The compressor turns it into CH₂ without loss; the "
            "liquefier into LH₂ at its 98% December yield. There is no LH₂ spot sale in the base "
            "case, so LH₂ storage has room for the extra RFNBO. Everything upstream (electrolyser, "
            "split output, H₂ buffer, pass-through) is as in test case 3.",
            False,
        ),
        ("", False),
        ("Proofs of sustainability (PoS sheet)", True),
        (
            "POS-A: dispatched 20 Oct, 8,000 kg, 3.0 gCO₂e/MJ (low GHG). POS-B: dispatched 20 "
            "Dec, 9,000 kg, 12.0 gCO₂e/MJ. POS-C: dispatched 25 Nov, 6,000 kg, 27.9 gCO₂e/MJ "
            "(near the 28.2 threshold). Without averaging (build plan §12) POS-C's part of either "
            "unit's batch would be above 28.2 (Parts!O13:O14) and no blend can save it, so all "
            "6,000 kg of it is not allocated and is written off at quarter end.",
            False,
        ),
        ("", False),
        ("Rules (build plan §11.1, decided 2026-10-17)", True),
        (
            "1. A PoS belongs to the quarter of its dispatch date, must be compliant (≤ 28.2), and "
            "its GHG includes distribution to the plant (no adder on receipt).",
            False,
        ),
        (
            "2. Its credentials are allocated to the supplier's metered delivery in the same "
            "quarter: into a unit per batch, at most the metered kg in that batch; per PoS at "
            "most its quantity. Credentials may go to flow metered before the dispatch date "
            "(decision 7): POS-B (20 Dec) covers delivery on 1–10 Nov.",
            False,
        ),
        (
            "3. Delivered H₂ without credentials is non-RFNBO, but it is still feed the unit "
            "consumed: RFNBO share = (RFNBO from storage + credentials) ÷ all feed. Credentials "
            "add their GHG to the feed emissions, like RFNBO feed from storage.",
            False,
        ),
        (
            "4. Storage weights count the supplier's delivery in the unit's feed: the H₂ buffer's "
            "weight is the liquefier's LH₂ ÷ (H₂ from the buffer + H₂ from the supplier).",
            False,
        ),
        (
            "5. PoS quantity not allocated by quarter end is written off (decision 2).",
            False,
        ),
        (
            "6. No averaging (§12, decided 2026-10-20): a batch fed from several sources has one "
            "part per source GHG and origin (Parts sheet), each keeping its source's GHG; "
            "purchased parts are never aggregated with the plant's own; a consignment's GHG is "
            "its highest part's.",
            False,
        ),
        ("", False),
        ("Why this decision set is optimal", True),
        (
            "Objective 2 (carried value): the compliant credentials (A and B, 17,000 kg) are all "
            "allocated: A 8,000 and B 1,600 kg into November's compressor delivery (9,600 kg), "
            "the rest of B (7,400 kg) into December's liquefier delivery. The 2,200 kg of "
            "December delivery no credential covers is filled with RFNBO from the H₂ buffer: "
            "December's liquefier yield (0.98) beats November's (0.97), so November's batch gives "
            "it up and holds 4,200 kg in the buffer, which ends full (4,880 kg). Every "
            "conditioning batch but November's liquefier batch is 100% RFNBO. LH₂ and CH₂ storage "
            "keep all their RFNBO after CA and CB, within measured inventory.",
            False,
        ),
        (
            "Objective 3 (carried emissions): CA and CB take their highest-GHG own parts within "
            "their limits first (the parts made from credentials are above both limits), so the "
            "lowest-GHG product stays in storage. Ties after that: each credential nearest its "
            "dispatch month: A (Oct) into November, B (Dec) into November's remainder and "
            "December.",
            False,
        ),
        ("", False),
        ("Variant (Variants sheet)", True),
        (
            "V1 — test case 3's 60,000 kg LH₂ spot sale is back. LH₂ storage now holds less than "
            "its RFNBO ledger, so a credential in the liquefier would only add LH₂ that is written "
            "off: none goes there (rule 5, PoS quantity is not pushed into product that would be "
            "written off). The compressor takes the lowest-GHG 9,600 kg; 13,400 kg of "
            "credentials are written off unallocated.",
            False,
        ),
        ("", False),
        ("Sheets", True),
        (
            "Parameters · PoS (register) · Data_15min · E_Batches · E_Split · Buffer_Ledger · "
            "PoS_Alloc (credentials to deliveries) · L_Batches · PT_Ledger · C_Batches · Parts · "
            "LH2_Ledger · CH2_Ledger · Consignments · Objectives · Variants · Checks.",
            False,
        ),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        s.put(f"A{i}", text, kind="text", bold=bold)
        s.ws[f"A{i}"].alignment = Alignment(wrap_text=True, vertical="top")


def parameters(wb: Workbook) -> None:
    s = Sheet(wb, "Parameters", {"A": 48, "B": 14, "C": 12, "D": 20})
    s.header(1, ["Parameter", "Value", "Unit", "Note"])
    for row, label, value, unit, note in PARAMS:
        s.note(f"A{row}", label)
        s.put(f"B{row}", value, fmt="General")
        s.note(f"C{row}", unit)
        s.note(f"D{row}", note)


def pos_sheet(wb: Workbook) -> None:
    s = Sheet(wb, "PoS", {"A": 12, "B": 18} | {get_column_letter(i): 15 for i in range(3, 15)})
    s.header(
        1,
        [
            "PoS",
            "Supplier",
            "Product",
            "Dispatch date",
            "Quantity (kg)",
            "GHG incl. distribution (gCO₂e/MJ)",
            "Mass balance period",
            "In this quarter?",
            "≤ 28.2 gCO₂e/MJ?",
            "Allocated (kg)",
            "Written off, not allocated (kg)",
            "Within quantity?",
            "Dispatch month",
        ],
    )
    for pid, r in POS_ROWS.items():
        p = T.POS[pid]
        s.put(f"A{r}", pid, kind="text")
        s.put(f"B{r}", T.SUPPLIER, kind="text")
        s.put(f"C{r}", "H2", kind="text")
        d: dt.date = p["dispatch"]
        s.put(f"D{r}", dt.datetime(d.year, d.month, d.day), kind="input", fmt="yyyy-mm-dd")
        s.put(f"E{r}", p["kg"], fmt="#,##0")
        s.put(f"F{r}", p["ghg"], fmt="0.0")
        s.put(f"G{r}", f'=YEAR(D{r})&"-Q"&ROUNDUP(MONTH(D{r})/3,0)', fmt="General")
        s.put(f"H{r}", f'=G{r}="2026-Q4"', fmt="General")
        s.put(f"I{r}", f"=F{r}<={P(8)}+0.000000001", fmt="General")
        s.put(f"J{r}", f"=SUMIF(PoS_Alloc!$A$2:$A$5,A{r},PoS_Alloc!$E$2:$E$5)")
        s.put(f"K{r}", f"=E{r}-J{r}", fill=KEY)
        s.put(f"L{r}", f"=J{r}<=E{r}+0.000001", fmt="General")
        s.put(f"M{r}", f'=YEAR(D{r})&"-"&TEXT(MONTH(D{r}),"00")', fmt="General")
    s.note("A5", "Total")
    for c in "EJK":
        s.put(f"{c}5", f"=SUM({c}2:{c}4)")
    s.note("A7", "Supplier reconciliation (H2 Trailers BV, H₂)")
    rec = [
        (8, "Metered delivery in the quarter (kg)", f"=SUM({rng('AC')})+SUM({rng('AD')})"),
        (9, "PoS quantity dispatched in the quarter (kg)", "=SUMPRODUCT(--(H2:H4),E2:E4)"),
        (10, "Credentials allocated (kg)", "=J5"),
        (11, "Delivered without credentials: non-RFNBO (kg)", "=B8-B10"),
        (12, "PoS quantity written off, not allocated (kg)", "=B9-B10"),
    ]
    for row, label, f in rec:
        s.note(f"A{row}", label)
        s.put(f"B{row}", f, fill=KEY if row == 12 else None)
    s.note(
        "A14",
        "A PoS is entered as received (blue). Not all the supplier's metered H₂ is RFNBO: only "
        "what credentials are allocated to. Credentials not allocated by quarter end are written "
        "off (decision 2).",
    )


def data_sheet(wb: Workbook, data: R.Data, d: R.Decisions) -> None:
    s = Sheet(wb, "Data_15min", {c: 14 for c, _, _ in DATA_COLS})
    s.ws.column_dimensions["B"].width = 18
    s.ws.column_dimensions["C"].width = 18
    s.header(1, [label for _, label, _ in DATA_COLS])
    for c in ("F", "G", "H", "Z", "AA", "AB"):
        s.ws[f"{c}1"].fill = DECISION
    s.ws.freeze_panes = "E2"
    ids = {u: [""] * len(data) for u in R.UNITS}
    for u in R.UNITS:
        for bid, i0, i1 in d.batches[u]:
            for i in range(i0, i1):
                ids[u][i] = bid
    blue = Font(name=FONT, size=10, color=W3.BLUE)
    plain = Font(name=FONT, size=10)
    ws = s.ws
    for i in range(len(data)):
        r = i + 2
        values: list[Any] = [
            i + 1,
            data.utc[i].replace(tzinfo=None),
            data.local[i].replace(tzinfo=None),
            data.month[i],
        ]
        for _, _, key in DATA_COLS[4:]:
            assert key is not None
            if key.startswith("alloc_"):
                values.append(d.renewable[key[6:]][i])
            elif key.startswith("batch_"):
                values.append(ids[key[6:]][i])
            else:
                values.append(data.s[key][i])
        for j, v in enumerate(values):
            cell = ws.cell(row=r, column=j + 1, value=v)
            cell.font = blue if isinstance(v, (int, float)) and j > 0 else plain
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=r, column=3).number_format = "yyyy-mm-dd hh:mm"


def buffer_ledger(wb: Workbook, title: str, *, pass_through: bool) -> None:
    """Test case 3's feed ledger. The buffer: October's and December's liquefier batches take
    all their buffer flow; November's leaves enough RFNBO to fill the buffer (its yield is
    below the buffer's weight)."""
    flow_col, target_sheet = ("E", "C_Batches") if pass_through else ("F", "L_Batches")
    prefix = "C" if pass_through else "L"
    targets = [f"{prefix}-2026-{m}" for m in ("10", "11", "12")]
    s = Sheet(wb, title, {"A": 44} | {get_column_letter(i): 16 for i in range(2, 9)})
    s.header(
        1,
        [
            "Electrolyser batch",
            "Batch GHG (gCO₂e/MJ)",
            "RFNBO H₂ in (kg)",
            *[f"→ {t} (kg)" for t in targets],
            "Left in storage (kg)",
            "Allocation within available?",
        ],
    )
    for k, r in enumerate((2, 3, 4)):
        s.put(f"A{r}", f"=E_Batches!A{r}")
        s.put(f"B{r}", f"=E_Batches!N{r}")
        s.put(f"C{r}", f"=E_Split!{flow_col}{r}")
        for j, col in enumerate("DEF"):
            # RFNBO from storage also covers the supplier's delivery that no credential covers.
            need = f"{target_sheet}!$T${r}+{target_sheet}!$U${r}-{target_sheet}!$W${r}"
            if j != k:
                s.put(f"{col}{r}", 0, fill=DECISION, fmt="#,##0.000000")
            elif pass_through or k != 1:
                s.put(f"{col}{r}", f"=MIN($C{r},{need})", fill=DECISION)
            else:
                s.put(f"{col}{r}", f"=MIN({need},$C{r}-MAX(0,$B$13-G2-G4))", fill=DECISION)
        s.put(f"G{r}", f"=C{r}-SUM(D{r}:F{r})")
        s.put(f"H{r}", f"=G{r}>=-0.000001", fmt="General")
    s.note("A5", "Total RFNBO H₂ (kg)")
    for c in "CDEFG":
        s.put(f"{c}5", f"=SUM({c}2:{c}4)")
    s.note("A6", "Feed GHG of each downstream batch (gCO₂e/MJ H₂)")
    for c in "DEF":
        s.put(f"{c}6", f"=SUMPRODUCT({c}2:{c}4,$B$2:$B$4)/{c}5")
    rows = [
        (9, "Opening RFNBO ledger (kg)", "=0" if pass_through else f"={P(10)}"),
        (10, "RFNBO H₂ in during the quarter (kg)", "=C5"),
        (11, "RFNBO H₂ fed downstream (kg)", "=SUM(D5:F5)"),
        (12, "Closing RFNBO ledger before write-off (kg)", "=B9+B10-B11"),
        (13, "Physical inventory at quarter end (kg)", "=0" if pass_through else f"={last('T')}"),
        (
            14,
            "Write-off (kg)",
            f"=MAX(0,B12-B13,SUMPRODUCT(--(B2:B4>{P(8)}+0.000000001),G2:G4))",
        ),
        (15, "RFNBO H₂ carried to Q1 2027 (kg)", "=B12-B14"),
        (
            16,
            "Energy-weighted GHG of carried H₂ (gCO₂e/MJ)",
            "=IF(B15>0,SUMPRODUCT(G2:G4,B2:B4)/B15,0)",
        ),
    ]
    for row, label, f in rows:
        s.note(f"A{row}", label)
        s.put(f"B{row}", f, fill=KEY if row == 15 else None)
    if pass_through:
        s.note("A17", "Pass-through ledger closes at zero?")
        s.put("B17", "=ABS(B15)<=0.000001", fmt="General")
        return
    s.note("A18", "Physical reconciliation")
    recon = [
        (19, "Opening physical inventory (kg)", f"={P(9)}"),
        (20, "H₂ in from the electrolyser (kg)", f"=SUM({rng('L')})"),
        (21, "H₂ consumed by the liquefier (kg)", f"=SUM({rng('O')})"),
        (22, "Expected closing (kg)", "=B19+B20-B21"),
        (23, "Measured closing (kg)", "=B13"),
        (24, "Implied loss (kg)", "=B22-B23"),
    ]
    for row, label, f in recon:
        s.note(f"A{row}", label)
        s.put(f"B{row}", f)
    s.note(
        "A26",
        "December's liquefier batch takes the 2,200 kg of supplier delivery no credential covers "
        "from the buffer as well (its yield, 0.98, beats the buffer's weight). November's takes "
        "less RFNBO than its buffer flow, so the buffer ends full of RFNBO (4,880 kg): 1 MJ of "
        "H₂ here is worth 0.9768 MJ of final product (Objectives!B3), more than the 0.97 MJ of "
        "LH₂ November's liquefier makes of it.",
    )


def pos_alloc(wb: Workbook) -> None:
    s = Sheet(wb, "PoS_Alloc", {"A": 22, "B": 14, "C": 14} | {c: 16 for c in "DEFGHI"})
    s.header(
        1,
        [
            "PoS",
            "Supply flow",
            "Batch",
            "Dispatch month",
            "Credentials allocated (kg)",
            "PoS GHG (gCO₂e/MJ)",
            "Emissions (gCO₂e)",
            "Batch month",
            "Months from dispatch",
        ],
    )
    for r, pid, flow, batch, f in ALLOC:
        look = f"MATCH(A{r},PoS!$A$2:$A$4,0)"
        s.put(f"A{r}", pid, kind="text", fill=DECISION)
        s.put(f"B{r}", flow, kind="text", fill=DECISION)
        s.put(f"C{r}", batch, kind="text", fill=DECISION)
        s.put(f"D{r}", f"=INDEX(PoS!$M$2:$M$4,{look})", fmt="General")
        s.put(f"E{r}", f, fill=DECISION)
        s.put(f"F{r}", f"=INDEX(PoS!$F$2:$F$4,{look})")
        s.put(f"G{r}", f"=E{r}*{P(17)}*F{r}")
        s.put(f"H{r}", f"=RIGHT(C{r},7)", fmt="General")
        s.put(
            f"I{r}",
            f"=ABS((LEFT(H{r},4)*12+RIGHT(H{r},2))-(LEFT(D{r},4)*12+RIGHT(D{r},2)))",
            fmt="0",
        )
    s.note("A6", "Total")
    for c in "EG":
        s.put(f"{c}6", f"=SUM({c}2:{c}5)")
    s.header(
        8,
        [
            "Delivery",
            "Supply flow",
            "Batch",
            "Metered delivery (kg)",
            "Credentials (kg)",
            "Within metered?",
            "First delivery (local)",
            "Covers delivery metered before the PoS's dispatch?",
        ],
    )
    for r, flow, batch, col, bcol in (
        (10, "f_sup_comp", "C-2026-11", "AD", "AB"),
        (11, "f_sup_liq", "L-2026-12", "AC", "AA"),
    ):
        s.put(f"A{r}", "compressor" if col == "AD" else "liquefier", kind="text")
        s.put(f"B{r}", flow, kind="text")
        s.put(f"C{r}", batch, kind="text")
        s.put(f"D{r}", f"=SUMIF({rng(bcol)},C{r},{rng(col)})")
        s.put(f"E{r}", f"=SUMIFS($E$2:$E$5,$B$2:$B$5,B{r},$C$2:$C$5,C{r})")
        s.put(f"F{r}", f"=E{r}<=D{r}+0.000001", fmt="General")
        s.put(
            f"G{r}",
            f'=_xlfn.MINIFS({rng("C")},{rng(bcol)},C{r},{rng(col)},">0")',
            fmt="yyyy-mm-dd hh:mm",
        )
    s.put(
        "H10",
        '=AND(SUMIFS(E2:E5,A2:A5,"POS-B",C2:C5,C10)>0,G10<=PoS!D3)',
        fmt="General",
    )
    s.put("H11", '=AND(SUMIFS(E2:E5,A2:A5,"POS-B",C2:C5,C11)>0,G11<=PoS!D3)', fmt="General")
    s.note("A12", "Total metered supply in the quarter (kg)")
    s.put("D12", f"=SUM({rng('AC')})+SUM({rng('AD')})")
    s.note("A13", "Metered delivery without credentials: non-RFNBO (kg)")
    s.put("D13", "=D12-E6")
    s.note(
        "A15",
        "Per batch and supply flow, credentials ≤ metered delivery; per PoS, ≤ its quantity "
        "(PoS!L). The compliant credentials are used in full (A and B); each goes to the "
        "delivery nearest its dispatch month (column I). POS-B (20 Dec) covers delivery metered "
        "before its dispatch date, within the quarter (decision 7). POS-C's part would be above "
        "28.2 in either unit (Parts!O13:O14) and cannot be blended down (§12), so none of it is "
        "allocated.",
    )


def unit_batches(
    wb: Workbook,
    title: str,
    *,
    ids: list[str],
    batch_col: str,
    consumed: str,
    supply: str,
    product: str,
    alloc: str,
    elec: str,
    feed_sheet: str,
    adder: int,
    product_label: str,
) -> None:
    """Test case 3's layout (columns A–S), with the feed from storage and from the supplier,
    and the credentials, in columns T–Y."""
    s = Sheet(wb, title, {"A": 14} | {get_column_letter(i): 15 for i in range(2, 26)})
    s.header(
        1,
        [
            "Batch ID",
            "Production month",
            "H₂ consumed, all feed (kg)",
            "RFNBO H₂ feed (kg)",
            "Feed GHG, all RFNBO feed (gCO₂e/MJ H₂; information, not averaged into GHG)",
            "RFNBO share",
            f"{product_label} produced (kg)",
            "Loss (kg)",
            "Loss (fraction of feed)",
            "Renewable allocated (kWh)",
            "Electricity used (kWh)",
            "Renewable matched (kWh)",
            "Own non-renewable electricity emissions (gCO₂e)",
            "Feed emissions (gCO₂e)",
            f"RFNBO {product_label} (kg)",
            f"RFNBO {product_label} (MJ)",
            "Highest part GHG (gCO₂e/MJ)",
            "Every part ≤ 28.2 gCO₂e/MJ?",
            "Intervals",
            "H₂ from storage (kg)",
            "H₂ from the supplier (kg)",
            "RFNBO from storage (kg)",
            "PoS credentials (kg)",
            "Emissions of RFNBO from storage (gCO₂e)",
            "Emissions of credentials (gCO₂e)",
        ],
    )
    B = rng(batch_col)
    for r, bid, fcol in zip((2, 3, 4), ids, "DEF", strict=True):
        s.put(f"A{r}", bid, kind="text", fill=DECISION)
        s.put(f"B{r}", f"=INDEX({rng('D')},MATCH(A{r},{B},0))")
        s.put(f"C{r}", f"=T{r}+U{r}")
        s.put(f"D{r}", f"=V{r}+W{r}")
        s.put(f"E{r}", f"=N{r}/(D{r}*{P(2)})")
        s.put(f"F{r}", f"=D{r}/C{r}")
        s.put(f"G{r}", f"=SUMIF({B},A{r},{rng(product)})")
        s.put(f"H{r}", f"=C{r}-G{r}")
        s.put(f"I{r}", f"=H{r}/C{r}")
        s.put(f"J{r}", f"=SUMIF({B},A{r},{rng(alloc)})")
        s.put(f"K{r}", f"=SUMIF({B},A{r},{rng(elec)})")
        s.put(f"L{r}", f"=MIN(J{r},K{r})")
        s.put(f"M{r}", f"=(K{r}-L{r})*{P(4)}")
        s.put(f"N{r}", f"=X{r}+Y{r}")
        s.put(f"O{r}", f"=F{r}*G{r}")
        s.put(f"P{r}", f"=O{r}*{P(2)}")
        s.put(f"Q{r}", f"=_xlfn.MAXIFS({PO},{PC},A{r})")
        s.put(
            f"R{r}",
            f"=SUMPRODUCT(({PC}=A{r})*({PO}>{P(8)}+0.000000001))=0",
            fmt="General",
        )
        s.put(f"S{r}", f"=COUNTIF({B},A{r})", fmt="#,##0")
        s.put(f"T{r}", f"=SUMIF({B},A{r},{rng(consumed)})")
        s.put(f"U{r}", f"=SUMIF({B},A{r},{rng(supply)})")
        s.put(f"V{r}", f"={feed_sheet}!{fcol}5")
        s.put(f"W{r}", f"=SUMIF(PoS_Alloc!$C$2:$C$5,A{r},PoS_Alloc!$E$2:$E$5)")
        s.put(f"X{r}", f"=V{r}*{P(2)}*{feed_sheet}!{fcol}6")
        s.put(f"Y{r}", f"=SUMIF(PoS_Alloc!$C$2:$C$5,A{r},PoS_Alloc!$G$2:$G$5)")
    s.note("A5", "Total")
    for c in "CDGHJKLMNOPSTUVWXY":
        s.put(f"{c}5", f"=SUM({c}2:{c}4)")
    s.note(
        "A7",
        "Share = (RFNBO from storage + credentials) ÷ all H₂ consumed, the supplier's H₂ "
        "included whether or not it carries credentials. No averaging (§12): the batch's RFNBO "
        "is split into parts, one per source GHG and origin (Parts sheet); each part keeps its "
        "source's GHG, and the batch is compliant when every part is.",
    )


def parts(wb: Workbook) -> None:
    """One row per part (§12): each keeps its source's GHG, adjusted by the unit's yield, own
    emissions and adder: GHG = (feed × LHV × g + share × own) ÷ (RFNBO product × LHV) + adder,
    which is g × C ÷ E + own ÷ E + adder."""
    s = Sheet(wb, "Parts", {"A": 34} | {get_column_letter(i): 15 for i in range(2, 17)})
    s.header(
        1,
        [
            "Part",
            "Unit",
            "Batch",
            "Production month",
            "Origin",
            "Source",
            "Source GHG (gCO₂e/MJ)",
            "RFNBO feed (kg)",
            "H₂ consumed by the batch (kg)",
            "Share",
            "Batch product (kg)",
            "RFNBO product (kg)",
            "Batch own emissions (gCO₂e)",
            "Adder (gCO₂e/MJ)",
            "Part GHG (gCO₂e/MJ)",
            "≤ 28.2 gCO₂e/MJ?",
        ],
    )
    for r, pid, unit, bid, origin, src, ghg, feed, br in PARTS:
        sheet, adder = ("L_Batches", 6) if unit == "liquefier" else ("C_Batches", 7)
        s.put(f"A{r}", pid, kind="text")
        s.put(f"B{r}", unit, kind="text")
        s.put(f"C{r}", bid, kind="text")
        s.put(f"D{r}", f"={sheet}!B{br}", fmt="General")
        s.put(f"E{r}", origin, kind="text")
        s.put(f"F{r}", src, kind="text")
        s.put(f"G{r}", f"={ghg}")
        s.put(f"H{r}", f"={feed}")
        s.put(f"I{r}", f"={sheet}!C{br}")
        s.put(f"J{r}", f"=H{r}/I{r}")
        s.put(f"K{r}", f"={sheet}!G{br}")
        s.put(f"L{r}", f"=J{r}*K{r}")
        s.put(f"M{r}", f"={sheet}!M{br}")
        s.put(f"N{r}", f"={P(adder)}")
        s.put(f"O{r}", f"=(H{r}*{P(2)}*G{r}+J{r}*M{r})/(L{r}*{P(2)})+N{r}", fill=KEY)
        s.put(f"P{r}", f"=O{r}<={P(8)}+0.000000001", fmt="General")
    s.note("A12", "Rejected: POS-C's part in either unit (its GHG does not depend on the kg)")
    for r, unit, bid, br in ((13, "compressor", "C-2026-11", 3), (14, "liquefier", "L-2026-12", 4)):
        sheet, adder = ("L_Batches", 6) if unit == "liquefier" else ("C_Batches", 7)
        s.put(f"A{r}", f"POS-C in {bid}", kind="text")
        s.put(f"B{r}", unit, kind="text")
        s.put(f"C{r}", bid, kind="text")
        s.put(f"F{r}", "POS-C", kind="text")
        s.put(f"G{r}", "=PoS!F4")
        s.put(f"I{r}", f"={sheet}!C{br}")
        s.put(f"K{r}", f"={sheet}!G{br}")
        s.put(f"M{r}", f"={sheet}!M{br}")
        s.put(f"N{r}", f"={P(adder)}")
        s.put(f"O{r}", f"=G{r}*I{r}/K{r}+M{r}/(K{r}*{P(2)})+N{r}")
        s.put(f"P{r}", f"=O{r}<={P(8)}+0.000000001", fmt="General")
    s.note("A15", "POS-C makes no compliant part?")
    s.put("B15", "=AND(NOT(P13),NOT(P14))", fmt="General")
    s.note(
        "A17",
        "No averaging (build plan §12, ISCC EU 203 §4.2): batches with different GHG are kept "
        "apart in the bookkeeping, so a unit batch fed from several sources has one part per "
        "source GHG and origin, named <batch>.1, .2, … in ascending GHG. Purchased parts are "
        "never aggregated with the plant's own. A part from a non-compliant source cannot be "
        "saved by blending.",
    )


def product_ledger(
    wb: Workbook,
    title: str,
    *,
    inv_col: str,
    opening_physical: int,
    opening_ledger: int,
    produced_col: str,
    shipped: list[tuple[str, str]],
) -> None:
    """Test case 3's product ledger, one line per part (§12), in write-off priority order."""
    order = LEDGER_ORDER[title]
    n, tr = len(order), total_row(title)
    s = Sheet(wb, title, {"A": 44} | {get_column_letter(i): 16 for i in range(2, 9)})
    s.header(
        1,
        [
            "Part (write-off priority order)",
            "Production month",
            "GHG (gCO₂e/MJ)",
            "In: RFNBO (kg)",
            "Out: consignment (kg)",
            "Closing before write-off (kg)",
            "Write-off (kg)",
            "Carried to Q1 2027 (kg)",
        ],
    )
    short = tr + 3
    for k, pid in enumerate(order):
        r, pr = k + 2, PART_ROW[pid]
        s.put(f"A{r}", pid, kind="text", fill=DECISION)
        s.put(f"B{r}", f"=Parts!D{pr}", fmt="General")
        s.put(f"C{r}", f"=Parts!O{pr}")
        s.put(f"D{r}", f"=Parts!L{pr}")
        s.put(f"E{r}", f"=SUMIF(Consignments!$D$6:$D$10,A{r},Consignments!$E$6:$E$10)")
        s.put(f"F{r}", f"=D{r}-E{r}")
        prev = "0" if k == 0 else f"SUM(G$2:G{r - 1})"
        s.put(f"G{r}", f"=MIN(F{r},MAX(0,$B${short}-{prev}))")
        s.put(f"H{r}", f"=F{r}-G{r}")
    s.note(f"A{tr}", "Total")
    for c in "DEFGH":
        s.put(f"{c}{tr}", f"=SUM({c}2:{c}{tr - 1})")
    lines = f"2:{{c}}{tr - 1}"
    s.note(f"A{tr + 2}", "Measured physical inventory at quarter end (kg)")
    s.put(f"B{tr + 2}", f"={last(inv_col)}")
    s.note(f"A{short}", "Shortfall to write off (kg): over inventory, and all non-compliant")
    s.put(
        f"B{short}",
        f"=MAX(0,{P(opening_ledger)}+F{tr}-B{tr + 2},"
        f"SUMPRODUCT(--(C{lines.format(c='C')}>{P(8)}+0.000000001),F{lines.format(c='F')}))",
    )
    s.note(f"A{tr + 4}", "Energy-weighted GHG of carried inventory (gCO₂e/MJ)")
    s.put(
        f"B{tr + 4}",
        f"=IF(H{tr}>0,SUMPRODUCT(C{lines.format(c='C')},H{lines.format(c='H')})/H{tr},0)",
    )
    s.note(f"A{tr + 5}", "Rows in priority order (highest GHG, then newest month)?")
    pair = "OR(ROUND(C{a},9)>ROUND(C{b},9),AND(ROUND(C{a},9)=ROUND(C{b},9),B{a}>B{b}))"
    s.put(
        f"B{tr + 5}",
        "=AND(" + ",".join(pair.format(a=r, b=r + 1) for r in range(2, 1 + n)) + ")",
        fmt="General",
    )
    s.note(f"A{tr + 7}", "Physical reconciliation")
    base = tr + 8
    rows = [
        (base, "Opening physical inventory (kg)", f"={P(opening_physical)}"),
        (base + 1, "Produced (kg)", f"=SUM({rng(produced_col)})"),
    ]
    for j, (label, col) in enumerate(shipped):
        rows.append((base + 2 + j, label, f"=SUM({rng(col)})"))
    m = base + 2 + len(shipped)
    out = "-".join(f"B{base + 2 + j}" for j in range(len(shipped)))
    rows += [
        (m, "Expected closing (kg)", f"=B{base}+B{base + 1}-{out}"),
        (m + 1, "Measured closing (kg)", f"=B{tr + 2}"),
        (m + 2, "Implied loss (kg)", f"=B{m}-B{m + 1}"),
    ]
    for row, label, f in rows:
        s.note(f"A{row}", label)
        s.put(f"B{row}", f)


def consignments(wb: Workbook) -> None:
    s = Sheet(wb, "Consignments", {"A": 30} | {get_column_letter(i): 15 for i in range(2, 19)})
    s.header(
        1,
        [
            "Consignment",
            "Customer",
            "Product",
            "Source storage",
            "Ship date",
            "Volume (kg)",
            "Target (gCO₂e/MJ)",
            "Distribution adder (gCO₂e/MJ)",
            "Limit for aggregated batches (gCO₂e/MJ)",
            "Aggregated batch: Oct 2026 (kg)",
            "Aggregated batch: Nov 2026 (kg)",
            "Aggregated batch: Dec 2026 (kg)",
            "Highest part GHG before shipping (gCO₂e/MJ)",
            "Shipped GHG (gCO₂e/MJ)",
            "Meets target?",
            "≤ 28.2 gCO₂e/MJ?",
            "Only compliant parts?",
            "Volume met?",
        ],
    )
    spec = [
        (2, "CA", "Customer A", "LH₂", "lh2_storage", 15, 150000, 5.0, 15),
        (3, "CB", "Customer B", "CH₂", "ch2_storage", 20, 120000, 4.0, 16),
    ]
    for r, cid, cust, prod, sto, day, vol, target, adder in spec:
        s.put(f"A{r}", cid, kind="text")
        s.put(f"B{r}", cust, kind="text")
        s.put(f"C{r}", prod, kind="text")
        s.put(f"D{r}", sto, kind="text")
        s.put(f"E{r}", dt.datetime(2026, 12, day), kind="input", fmt="yyyy-mm-dd")
        s.put(f"F{r}", vol, fmt="#,##0")
        s.put(f"G{r}", target, fmt="0.0")
        s.put(f"H{r}", f"={P(adder)}")
        s.put(f"I{r}", f"=G{r}-H{r}")
        months = [("J", "2026-10"), ("K", "2026-11"), ("L", "2026-12")]
        for col, month in months:
            s.put(f"{col}{r}", f'=SUMIFS($E$6:$E$10,$B$6:$B$10,A{r},$C$6:$C$10,"{month}")')
        # No average (§12 decision 4): every part plus the adder meets the target.
        s.put(f"M{r}", f"=_xlfn.MAXIFS($F$6:$F$10,$B$6:$B$10,A{r})")
        s.put(f"N{r}", f"=M{r}+H{r}", fill=KEY)
        s.put(f"O{r}", f"=N{r}<=G{r}+0.000000001", fmt="General")
        s.put(f"P{r}", f"=N{r}<={P(8)}+0.000000001", fmt="General")
        s.put(
            f"Q{r}",
            f"=SUMPRODUCT(($B$6:$B$10=A{r})*($F$6:$F$10>{P(8)}+0.000000001))=0",
            fmt="General",
        )
        s.put(f"R{r}", f"=ABS(J{r}+K{r}+L{r}-F{r})<=0.000001*F{r}", fmt="General")
    s.header(
        5,
        [
            "Aggregated batch",
            "Consignment",
            "Production month",
            "Part",
            "kg",
            "Part GHG",
            "≤ 28.2?",
        ],
    )
    # Highest GHG first among the parts within each consignment's limit (I2, I3), then the
    # oldest month. The purchased parts (3.5 and above) are above both limits.
    agg = [
        (6, "CA-2026-11", "CA", "2026-11", "L-2026-11", f"=MIN(F2,Parts!L{PART_ROW['L-2026-11']})"),
        (7, "CA-2026-10", "CA", "2026-10", "L-2026-10", "=F2-E6"),
        (8, "CB-2026-10", "CB", "2026-10", "C-2026-10", f"=MIN(F3,Parts!L{PART_ROW['C-2026-10']})"),
        (
            9,
            "CB-2026-11",
            "CB",
            "2026-11",
            "C-2026-11.1",
            f"=MIN(F3-E8,Parts!L{PART_ROW['C-2026-11.1']})",
        ),
        (10, "CB-2026-12", "CB", "2026-12", "C-2026-12", "=F3-E8-E9"),
    ]
    for r, aid, cid, month, src, f in agg:
        s.put(f"A{r}", aid, kind="text")
        s.put(f"B{r}", cid, kind="text")
        s.put(f"C{r}", month, kind="text")
        s.put(f"D{r}", src, kind="text")
        s.put(f"E{r}", f, fill=DECISION)
        s.put(f"F{r}", f"=INDEX({PO},MATCH(D{r},{PA},0))")
        s.put(f"G{r}", f"=F{r}<={P(8)}+0.000000001", fmt="General")
    s.note(
        "A12",
        "No averaging (§12): each part of a consignment plus the adder must meet the target, so "
        "the consignment's GHG is its highest part's. The parts made from credentials (CH₂ 3.5 "
        "and 12.5, LH₂ 12.7) are above CA's and CB's limits; each consignment takes its "
        "highest-GHG own parts within the limit first, then the oldest month, so the lowest-GHG "
        "product stays in storage.",
    )


def objectives(wb: Workbook) -> None:
    s = Sheet(wb, "Objectives", {"A": 70, "B": 18, "C": 70})
    s.header(1, ["Objective (in order)", "Value", "Note"])
    lhv = P(2)
    tl, tc = total_row("LH2_Ledger"), total_row("CH2_Ledger")
    rows: list[tuple[int, str, Any, str]] = [
        (
            2,
            "1. Consignments met",
            "=AND(Consignments!O2,Consignments!P2,Consignments!Q2,Consignments!R2)"
            "+AND(Consignments!O3,Consignments!P3,Consignments!Q3,Consignments!R3)",
            "Of 2",
        ),
        (
            3,
            "Storage weight: H₂ buffer (MJ final product per MJ)",
            f"=SUM({rng('P')})/(SUM({rng('O')})+SUM({rng('AC')}))",
            "Liquefier energy yield over the quarter, the supplier's H₂ in its feed",
        ),
        (
            4,
            "Storage weight: pass-through",
            f"=SUM({rng('S')})/(SUM({rng('R')})+SUM({rng('AD')}))",
            "Compressor energy yield; the pass-through always closes at zero",
        ),
        (5, "Storage weight: LH₂ storage", 1, "Final product"),
        (6, "Storage weight: CH₂ storage", 1, "Final product"),
        (7, "Compliant RFNBO carried: H₂ buffer (MJ)", f"=Buffer_Ledger!B15*{lhv}", ""),
        (8, "Compliant RFNBO carried: pass-through (MJ)", f"=PT_Ledger!B15*{lhv}", ""),
        (9, "Compliant RFNBO carried: LH₂ storage (MJ)", f"=LH2_Ledger!H{tl}*{lhv}", ""),
        (10, "Compliant RFNBO carried: CH₂ storage (MJ)", f"=CH2_Ledger!H{tc}*{lhv}", ""),
        (
            11,
            "2. Compliant RFNBO carried forward (MJ of final product)",
            "=SUMPRODUCT(B3:B6,B7:B10)",
            "Key answer: each storage's carried MJ times its weight",
        ),
        (
            12,
            "3. Carried emissions (gCO₂e)",
            f"=Buffer_Ledger!B15*Buffer_Ledger!B16*{lhv}"
            f"+SUMPRODUCT(LH2_Ledger!C2:C{tl - 1},LH2_Ledger!H2:H{tl - 1})*{lhv}"
            f"+SUMPRODUCT(CH2_Ledger!C2:C{tc - 1},CH2_Ledger!H2:H{tc - 1})*{lhv}",
            "",
        ),
        (
            13,
            "4. Batches",
            "=COUNTA(E_Batches!A2:A4)+COUNTA(L_Batches!A2:A4)+COUNTA(C_Batches!A2:A4)",
            "One per unit and month",
        ),
        (
            14,
            "Tie-break: purchased emissions allocated (gCO₂e)",
            "=PoS_Alloc!G6",
            "The lowest-GHG credentials where the objectives tie",
        ),
        (15, "PoS quantity written off, not allocated (kg)", "=PoS!K5", "Reported"),
    ]
    for row, label, f, note in rows:
        s.note(f"A{row}", label)
        s.put(f"B{row}", f, fill=KEY if row in (2, 11, 15) else None)
        s.note(f"C{row}", note)
    for row in (2, 13):
        s.ws[f"B{row}"].number_format = "0"
    s.note("A17", "Upper-bound argument for objective 2")
    bound = [
        (18, "Every electrolyser batch 100% RFNBO?", "=MIN(E_Batches!I2:I4)>=1-0.000000001"),
        (
            19,
            "Every compliant credential allocated (POS-C makes no compliant part)?",
            "=AND(PoS!K2=0,PoS!K3=0,Parts!B15)",
        ),
        (
            20,
            "H₂ buffer holds RFNBO up to its measured inventory?",
            "=ABS(Buffer_Ledger!B15-Buffer_Ledger!B13)<=0.000001",
        ),
        (
            21,
            "November liquefier yield below the buffer's weight (so the buffer stays full)?",
            "=L_Batches!G3/L_Batches!C3<B3",
        ),
        (
            22,
            "October and December liquefier batches fully RFNBO (yield above the weight)?",
            "=AND(L_Batches!F2>=1-0.000000001,L_Batches!F4>=1-0.000000001,"
            "L_Batches!G2/L_Batches!C2>B3,L_Batches!G4/L_Batches!C4>B3)",
        ),
        (
            23,
            "LH₂ and CH₂ storage hold all their RFNBO less CA and CB, within inventory?",
            f"=AND(ABS(LH2_Ledger!H{tl}-(L_Batches!O5-Consignments!F2))<=0.000001,"
            f"ABS(CH2_Ledger!H{tc}-(C_Batches!O5-Consignments!F3))<=0.000001,"
            f"LH2_Ledger!G{tl}=0,CH2_Ledger!G{tc}=0)",
        ),
        (24, "Objective 2 at its upper bound?", "=AND(B18:B23)"),
    ]
    for row, label, f in bound:
        s.note(f"A{row}", label)
        s.put(f"B{row}", f, fmt="General")
    s.note(
        "A26",
        "Credentials add to objective 2 only where the product they make is kept: here every "
        "part they make stays within its storage's inventory. Variant V1 shows the opposite. "
        "December's liquefier delivery that no compliant credential covers (2,200 kg) is filled "
        "with RFNBO from the buffer, at December's 0.98 yield rather than November's 0.97.",
    )


def variants(wb: Workbook) -> None:
    s = Sheet(wb, "Variants", {"A": 74, "B": 18, "C": 90})
    s.header(1, ["Variant", "Value", "Note"])
    lhv, w = P(2), "Objectives!$B$3"
    tl, tc = total_row("LH2_Ledger"), total_row("CH2_Ledger")
    rows: list[tuple[int, str, Any, str]] = [
        (2, "V1 — LH₂ spot sale back (no RFNBO claim)", None, "Test case 3's 60,000 kg on 29 Dec"),
        (3, "Spot sale (kg)", R.SPOT_KG, "Input"),
        (4, "LH₂ storage measured inventory at quarter end (kg)", f"=LH2_Ledger!B{tl + 2}-B3", ""),
        (
            5,
            "H₂ the December liquefier batch takes from the buffer (kg)",
            "=Buffer_Ledger!C4-Buffer_Ledger!B13",
            "Test case 3's routing: December's batch fills the buffer; November takes all its flow",
        ),
        (
            6,
            "December LH₂ RFNBO, no credentials (kg)",
            "=B5/L_Batches!C4*L_Batches!G4",
            "Share = buffer RFNBO ÷ all feed (supplier's H₂ is non-RFNBO)",
        ),
        (
            7,
            "LH₂ RFNBO ledger before write-off (kg)",
            "=L_Batches!O2+L_Batches!G3+B6-Consignments!F2",
            "October and November fully RFNBO; CA as before",
        ),
        (8, "LH₂ written off: ledger over inventory (kg)", "=MAX(0,B7-B4)", ""),
        (9, "LH₂ carried (kg)", "=B7-B8", "Equals the measured inventory"),
        (
            10,
            "LH₂ carried if December's batch also took 9,600 kg of credentials (kg)",
            "=MIN(B7+L_Batches!U4/L_Batches!C4*L_Batches!G4,B4)",
            "Still the measured inventory: the extra LH₂ would be written off",
        ),
        (11, "Credentials into the liquefier (kg)", 0, "Expected decision: none"),
        (
            12,
            "Credentials into the compressor (kg)",
            "=MIN(PoS_Alloc!D10,PoS!E2+PoS!E3)",
            "Lowest GHG first: all of POS-A, then POS-B",
        ),
        (13, "PoS quantity written off, not allocated (kg)", "=PoS!E5-B11-B12", ""),
        (14, "H₂ buffer carried (kg)", "=Buffer_Ledger!B13", "Full"),
        (15, "CH₂ carried (kg)", f"=CH2_Ledger!H{tc}", "As the base case: 9,600 kg of credentials"),
        (
            16,
            "RFNBO carried forward (MJ final product)",
            f"=(B14*{w}+B9+B15)*{lhv}",
            "Key answer",
        ),
        (
            17,
            "Credentials in the liquefier would add nothing?",
            "=ABS(B10-B9)<=0.000001",
            "Rule 5: PoS quantity is not pushed into product that would be written off",
        ),
    ]
    for row, label, f, note in rows:
        s.note(f"A{row}", label)
        if f is not None:
            s.put(f"B{row}", f, fill=KEY if row in (13, 16) else None)
        s.note(f"C{row}", note)
    s.ws["B17"].number_format = "General"
    s.ws["A2"].font = Font(name=FONT, size=10, bold=True)


def check_rows(cases: dict[str, Any]) -> list[tuple[str, str, float]]:
    b = cases["base"]
    u, led, cons, obj, pos = b["units"], b["ledgers"], b["consignments"], b["objectives"], b["pos"]
    out: list[tuple[str, str, float]] = [
        ("Intervals in quarter", f"=COUNT({rng('A')})", b["intervals"]),
        ("Metered supply (kg)", "=PoS!B8", b["metered_supply_kg"]),
    ]
    for pid, r in POS_ROWS.items():
        out += [
            (f"{pid} allocated (kg)", f"=PoS!J{r}", pos[pid]["allocated_kg"]),
            (f"{pid} written off (kg)", f"=PoS!K{r}", pos[pid]["written_off_kg"]),
        ]
    for unit, sheet, ids in (
        ("liquefier", "L_Batches", ("L-2026-10", "L-2026-11", "L-2026-12")),
        ("compressor", "C_Batches", ("C-2026-10", "C-2026-11", "C-2026-12")),
    ):
        for r, bid in zip((2, 3, 4), ids, strict=True):
            x = u[unit][bid]
            out += [
                (f"{bid} feed consumed (kg)", f"={sheet}!C{r}", x["feed_consumed_kg"]),
                (f"{bid} credentials (kg)", f"={sheet}!W{r}", x["pos_feed_kg"]),
                (f"{bid} RFNBO feed (kg)", f"={sheet}!D{r}", x["rfnbo_feed_kg"]),
                (f"{bid} RFNBO share", f"={sheet}!F{r}", x["rfnbo_share"]),
                (f"{bid} RFNBO product (kg)", f"={sheet}!O{r}", x["rfnbo_kg"]),
                (
                    f"{bid} highest part GHG (gCO₂e/MJ)",
                    f"={sheet}!Q{r}",
                    max(y["ghg"] for y in x["parts"].values()),
                ),
            ]
            for pid, y in x["parts"].items():
                pr = PART_ROW[pid]
                out += [
                    (f"{pid} RFNBO (kg)", f"=Parts!L{pr}", y["rfnbo_kg"]),
                    (f"{pid} GHG (gCO₂e/MJ)", f"=Parts!O{pr}", y["ghg"]),
                ]
    out += [
        ("H₂ buffer carried (kg)", "=Buffer_Ledger!B15", led["h2_buffer"]["carried_kg"]),
        (
            "Pass-through closing ledger (kg)",
            "=PT_Ledger!B12",
            led["pt_compressor"]["closing_before_kg"],
        ),
    ]
    for sheet, storage in (("LH2_Ledger", "lh2_storage"), ("CH2_Ledger", "ch2_storage")):
        lines = led[storage]["lines"]
        assert set(lines) == set(LEDGER_ORDER[sheet]), sheet
        for r, pid in enumerate(LEDGER_ORDER[sheet], 2):
            out += [
                (f"{storage} write-off {pid} (kg)", f"={sheet}!G{r}", lines[pid]["write_off_kg"]),
                (f"{storage} carried {pid} (kg)", f"={sheet}!H{r}", lines[pid]["carried_kg"]),
            ]
        out.append(
            (
                f"{storage} measured closing (kg)",
                f"={sheet}!B{total_row(sheet) + 2}",
                led[storage]["physical_kg"],
            )
        )
    for r, cid in ((2, "CA"), (3, "CB")):
        out.append(
            (f"{cid} shipped GHG (gCO₂e/MJ)", f"=Consignments!N{r}", cons[cid]["shipped_ghg"])
        )
    out += [
        ("Consignments met", "=Objectives!B2", obj["consignments_met"]),
        ("H₂ buffer weight", "=Objectives!B3", obj["weights"]["h2_buffer"]),
        ("RFNBO carried forward (MJ final product)", "=Objectives!B11", obj["carried_value_MJ"]),
        ("Carried emissions (gCO₂e)", "=Objectives!B12", obj["carried_emissions_g"]),
        ("Batches", "=Objectives!B13", obj["batches"]),
        ("Purchased emissions allocated (gCO₂e)", "=Objectives!B14", obj["purchased_emissions_g"]),
    ]
    v1 = cases["v1_spot"]
    out += [
        (
            "V1 December LH₂ RFNBO (kg)",
            "=Variants!B6",
            v1["units"]["liquefier"]["L-2026-12"]["rfnbo_kg"],
        ),
        ("V1 LH₂ written off (kg)", "=Variants!B8", v1["ledgers"]["lh2_storage"]["write_off_kg"]),
        ("V1 LH₂ carried (kg)", "=Variants!B9", v1["ledgers"]["lh2_storage"]["carried_kg"]),
        (
            "V1 credentials into the compressor (kg)",
            "=Variants!B12",
            v1["units"]["compressor"]["C-2026-11"]["pos_feed_kg"],
        ),
        (
            "V1 PoS written off (kg)",
            "=Variants!B13",
            sum(p["written_off_kg"] for p in v1["pos"].values()),
        ),
        ("V1 RFNBO carried (MJ)", "=Variants!B16", v1["objectives"]["carried_value_MJ"]),
    ]
    return out


def checks(wb: Workbook, cases: dict[str, Any]) -> int:
    s = Sheet(wb, "Checks", {"A": 52, "B": 20, "C": 26, "D": 14, "E": 8})
    s.header(1, ["Check", "Workbook", "Independent Python result", "Difference", "Pass?"])
    rows = check_rows(cases)
    for i, (label, f, value) in enumerate(rows, 2):
        s.note(f"A{i}", label)
        s.put(f"B{i}", f)
        s.put(f"C{i}", float(value))
        s.put(f"D{i}", f"=B{i}-C{i}")
        s.put(f"E{i}", f"=ABS(D{i})<=0.000001*MAX(1,ABS(C{i}))", fmt="General")
    end = len(rows) + 1
    flags = [
        "E_Split!G2:G4",
        "E_Split!H2:H4",
        "Buffer_Ledger!H2:H4",
        "PT_Ledger!H2:H4",
        "PT_Ledger!B17",
        f"LH2_Ledger!B{total_row('LH2_Ledger') + 5}",
        f"CH2_Ledger!B{total_row('CH2_Ledger') + 5}",
        "Parts!B15",
        "PoS!H2:H4",
        "PoS!I2:I4",
        "PoS!L2:L4",
        "PoS_Alloc!F10:F11",
        "PoS_Alloc!H10:H11",
        "Objectives!B24",
        "Variants!B17",
    ]
    s.note(
        f"A{end + 2}",
        "Rule flags all TRUE (split, allocations, write-off order, pass-through, PoS, bound)?",
    )
    s.put(f"E{end + 2}", f"=AND({','.join(flags)})", fmt="General")
    s.note(f"A{end + 3}", "Renewable pool never over-allocated in any interval?")
    s.put(
        f"E{end + 3}",
        f"=SUMPRODUCT(--({rng('F')}+{rng('G')}+{rng('H')}>{rng('E')}+0.000000001))=0",
        fmt="General",
    )
    s.note(f"A{end + 5}", "All checks pass?")
    s.put(
        f"E{end + 5}",
        f"=AND(COUNTIF(E2:E{end},FALSE())=0,E{end + 2},E{end + 3})",
        fill=KEY,
        fmt="General",
    )
    s.note(
        f"A{end + 6}",
        "Column C was computed by tests/golden/tc4_reference.py from the same inputs, not from "
        "these formulas.",
    )
    return end + 5


def build(out: Path) -> Path:
    data = T.build_data()
    d = T.base_decisions(data).base
    cases = T.all_cases()
    wb = Workbook()
    wb.remove(wb.active)
    readme(wb)
    parameters(wb)
    pos_sheet(wb)
    data_sheet(wb, data, d)
    W3.e_batches(wb)
    W3.e_split(wb)
    buffer_ledger(wb, "Buffer_Ledger", pass_through=False)
    pos_alloc(wb)
    unit_batches(
        wb,
        "L_Batches",
        ids=["L-2026-10", "L-2026-11", "L-2026-12"],
        batch_col="AA",
        consumed="O",
        supply="AC",
        product="P",
        alloc="G",
        elec="N",
        feed_sheet="Buffer_Ledger",
        adder=6,
        product_label="LH₂",
    )
    buffer_ledger(wb, "PT_Ledger", pass_through=True)
    unit_batches(
        wb,
        "C_Batches",
        ids=["C-2026-10", "C-2026-11", "C-2026-12"],
        batch_col="AB",
        consumed="R",
        supply="AD",
        product="S",
        alloc="H",
        elec="Q",
        feed_sheet="PT_Ledger",
        adder=7,
        product_label="CH₂",
    )
    parts(wb)
    product_ledger(
        wb,
        "LH2_Ledger",
        inv_col="W",
        opening_physical=11,
        opening_ledger=12,
        produced_col="P",
        shipped=[("Shipped CA (kg)", "U"), ("Spot sale, no claim (kg)", "V")],
    )
    product_ledger(
        wb,
        "CH2_Ledger",
        inv_col="Y",
        opening_physical=13,
        opening_ledger=14,
        produced_col="S",
        shipped=[("Shipped CB (kg)", "X")],
    )
    consignments(wb)
    objectives(wb)
    variants(wb)
    checks(wb, cases)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("out", nargs="?", type=Path, default=OUT)
    print(build(ap.parse_args().out))


if __name__ == "__main__":
    main()
